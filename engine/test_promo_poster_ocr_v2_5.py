"""Regression tests for v2.5 context taxonomy and cross-source review signals."""
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))
import promo_poster_ocr as ocr


def _line(text, x0, y0, x1, y1, confidence=0.97):
    return {
        "text": text,
        "confidence": confidence,
        "x0": x0,
        "y0": y0,
        "x1": x1,
        "y1": y1,
    }


def test_member_discount_is_not_misreported_as_member_product_price():
    result = ocr.classify_promo_context("Khusus Member Potongan Rp 500")
    assert result["member_offer_classification"] == "POTONGAN_MEMBER"
    assert result["member_discount_candidate"] == 500
    assert result["member_price_candidate"] == ""
    assert ocr._clean_local_product_line("Khusus Member Potongan Rp 500") == ""


def test_minimum_spend_is_not_misreported_as_member_product_price():
    result = ocr.classify_promo_context("Member Alfamidi minimal Rp100.000")
    assert result["member_offer_classification"] == "MINIMUM_BELANJA"
    assert result["minimum_spend_candidate"] == 100000
    assert result["member_price_candidate"] == ""


def test_regional_member_price_list_is_separate_from_one_member_price():
    text = "Member: Sumatera Rp39.500, Kalimantan Rp41.000"
    result = ocr.classify_promo_context(text)
    assert result["member_offer_classification"] == "HARGA_REGIONAL"
    assert result["member_price_candidate"] == ""
    assert "Sumatera" in result["regional_price_context"]
    assert "Kalimantan" in result["regional_price_context"]
    assert "39.500" in result["regional_price_context"]
    assert ocr._clean_local_product_line(text) == ""


def test_standalone_legal_footer_is_not_a_product_name():
    assert ocr._is_campaign_legal_line("Syarat ketentuan s.d Oktober Hanya di Bayar Tagihan")
    assert ocr._clean_local_product_line("Syarat ketentuan s.d Oktober Hanya di Bayar Tagihan") == ""
    assert ocr._is_campaign_legal_line("Produk dapat ditebus maksimal 14 hari sejak pemesanan PROMISE PASTI dibuat")
    assert ocr._clean_local_product_line("Produk dapat ditebus maksimal 14 hari sejak pemesanan PROMISE PASTI dibuat") == ""


def test_leading_city_list_is_removed_but_product_title_is_retained():
    title = "Kupang Ternate Jayapura Abepura Wafer double ckt matcha ASTOR"
    cleaned = ocr._clean_local_product_line(title)
    assert cleaned == "Wafer double ckt matcha ASTOR"
    assert not any(city in cleaned.casefold() for city in ("kupang", "ternate", "jayapura", "abepura"))


def test_regional_restriction_suffix_is_removed_from_product_title():
    title = "Mie Seirocka Tori Cup 114g Tidak berlaku di Ambon"
    assert ocr._clean_local_product_line(title) == "Mie Seirocka Tori Cup 114g"
    result = ocr.classify_promo_context(title)
    assert "Ambon" in result["regional_scope"]


def test_cross_source_price_variation_is_review_signal_not_auto_decision():
    candidates = [
        {
            "candidate_id": "A", "article_url": "https://example.test/a", "image_url": "https://example.test/a.jpg",
            "campaign_title": "Promo Hemat 26-31 Oktober", "matched_product_name": "Vinda Deluxe Kitchen Towel 2 Roll",
            "normal_price_candidate": 47900, "promo_price_candidate": 30500,
            "review_priority": "RENDAH", "review_priority_reason": "Belum ada pemicu.",
        },
        {
            "candidate_id": "B", "article_url": "https://other.test/a", "image_url": "https://other.test/b.jpg",
            "campaign_title": "Promo Hemat 26-31 Oktober", "matched_product_name": "Vinda Deluxe Kitchen Towel 2 Roll",
            "normal_price_candidate": 47900, "promo_price_candidate": 31900,
            "review_priority": "RENDAH", "review_priority_reason": "Belum ada pemicu.",
        },
    ]
    result = ocr._annotate_cross_source_price_variations(candidates)
    assert all(row["cross_source_price_variation_status"] == "VARIASI_HARGA_PERLU_VERIFIKASI" for row in result)
    assert all("periksa periode" in row["cross_source_price_variation_note"].casefold() for row in result)
    assert all(row["review_priority"] in {"SEDANG", "TINGGI"} for row in result)
    # The comparison cannot approve any candidate.
    assert all(row.get("decision", "CEK") == "CEK" for row in result)


def test_cross_source_comparison_requires_same_campaign_title_and_article_distinction():
    candidates = [
        {"candidate_id": "A", "article_url": "https://example.test/a", "image_url": "a.jpg", "campaign_title": "Promo Minggu 1", "matched_product_name": "Aqua 600ml", "normal_price_candidate": 11900, "promo_price_candidate": 10900},
        {"candidate_id": "B", "article_url": "https://other.test/b", "image_url": "b.jpg", "campaign_title": "Promo Minggu 2", "matched_product_name": "Aqua 600ml", "normal_price_candidate": 11900, "promo_price_candidate": 9900},
    ]
    result = ocr._annotate_cross_source_price_variations(candidates)
    assert all(not row.get("cross_source_price_variation_status") for row in result)


def test_cross_source_comparison_skips_quarantined_local_names():
    candidates = [
        {"candidate_id": "A", "article_url": "https://example.test/a", "image_url": "a.jpg", "campaign_title": "Promo Sama", "local_product_name_suggestion": "GoPay Bayar Pakai", "local_product_name_suggestion_risk": "TINGGI", "local_product_name_suggestion_status": "DIKARANTINA", "normal_price_candidate": 11900, "promo_price_candidate": 10900},
        {"candidate_id": "B", "article_url": "https://other.test/b", "image_url": "b.jpg", "campaign_title": "Promo Sama", "local_product_name_suggestion": "GoPay Bayar Pakai", "local_product_name_suggestion_risk": "TINGGI", "local_product_name_suggestion_status": "DIKARANTINA", "normal_price_candidate": 11900, "promo_price_candidate": 9900},
    ]
    result = ocr._annotate_cross_source_price_variations(candidates)
    assert all(not row.get("cross_source_price_variation_status") for row in result)


def test_review_workbook_places_v25_fields_under_the_correct_headers(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "review.xlsx"
    candidate = {
        "candidate_id": "TESTV25", "normal_price_candidate": 100000, "promo_price_candidate": 90000,
        "member_price_candidate": "", "member_discount_candidate": 500, "minimum_spend_candidate": 100000,
        "regional_price_context": "Member Sumatera Rp39.500, Kalimantan Rp41.000",
        "member_offer_classification": "POTONGAN_MEMBER",
        "cross_source_price_variation_status": "VARIASI_HARGA_PERLU_VERIFIKASI",
        "cross_source_price_variation_note": "Periksa periode dan wilayah.",
        "local_product_name_suggestion_status": "SARAN OCR LOKAL — PERLU VERIFIKASI",
        "review_priority": "SEDANG", "review_priority_reason": "Uji", "ocr_snippet": "test",
    }
    ocr.make_review_workbook([candidate], path)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Kandidat OCR"]
    headers = [cell.value for cell in ws[1]]
    row = [cell.value for cell in ws[2]]
    assert len(headers) == len(row) == 47
    values = dict(zip(headers, row))
    assert values["Potongan Member (Rp)"] == 500
    assert values["Minimum Belanja (Rp)"] == 100000
    assert values["Klasifikasi Harga/Ketentuan Member"] == "POTONGAN_MEMBER"
    assert values["Variasi Harga Antar-Sumber"] == "VARIASI_HARGA_PERLU_VERIFIKASI"
    assert values["Catatan Perbandingan Harga Antar-Sumber"] == "Periksa periode dan wilayah."
    wb.close()
