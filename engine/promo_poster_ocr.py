"""Local OCR for promotion-catalog poster images.

No generative price guessing is used. OCR lines and candidate price pairs are
saved with image/page provenance; only rows explicitly approved in Excel are
eligible to fill missing promo/normal price values in the final catalog.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import ipaddress
import json
import re
import statistics
import sys
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse, urlunparse

import requests
from PIL import Image, ImageEnhance, ImageOps, UnidentifiedImageError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from promo_names import product_name_similarity, product_name_size_signature
from promo_price_reference import identity_key
from promo_market_prices import parse_idr_price
from promo_poster_overrides import DEFAULT_REVIEW_XLSX, import_reviewed_prices

ROOT = Path(__file__).resolve().parent.parent
IMAGE_INDEX = ROOT / "_internal" / "promo_multisource_catalog_images.csv"
PRODUCTS_CSV = ROOT / "_internal" / "promo_multisource_products_all.csv"
CACHE_CSV = ROOT / "_internal" / "promo_poster_ocr_cache.csv"
LINES_CSV = ROOT / "_internal" / "promo_poster_ocr_lines.csv"
CANDIDATES_CSV = ROOT / "_internal" / "promo_poster_price_candidates.csv"
REVIEW_XLSX = DEFAULT_REVIEW_XLSX
OVERRIDES_CSV = ROOT / "_internal" / "promo_poster_price_overrides.csv"
MAX_IMAGE_BYTES = 20 * 1024 * 1024
REQUEST_TIMEOUT = (5, 15)
USER_AGENT = "Mozilla/5.0 (compatible; MerchantsPromoPosterOCR/1.0; local catalog audit)"

CACHE_FIELDS = [
    "image_url", "source_name", "campaign_title", "article_url", "image_alt", "period_status",
    "scraped_at", "image_sha256", "image_width", "image_height", "ocr_status", "ocr_backend",
    "ocr_error", "processed_at", "lines_json",
]
LINE_FIELDS = [
    "source_name", "campaign_title", "article_url", "image_url", "image_alt", "period_status",
    "line_no", "ocr_text", "ocr_confidence", "x0", "y0", "x1", "y1", "image_width", "image_height", "processed_at",
]
CANDIDATE_FIELDS = [
    "candidate_id", "source_name", "campaign_title", "article_url", "image_url", "image_alt",
    "ocr_snippet", "price_values_json", "normal_price_candidate", "promo_price_candidate",
    "discount_claim_percent", "discount_calculated_percent", "ocr_confidence",
    "matched_product_name", "matched_product_score", "match_note", "candidate_note",
]
REVIEW_HEADERS = [
    "ID Kandidat", "Keputusan", "Nama Produk Disetujui", "Nama Produk Saran",
    "Harga Normal Disetujui (Rp)", "Harga Promo Disetujui (Rp)",
    "Harga Normal Kandidat (Rp)", "Harga Promo Kandidat (Rp)",
    "Klaim Diskon OCR (%)", "Diskon Hitung (%)", "Kecocokan Nama", "Keyakinan OCR",
    "Potongan Teks Poster OCR", "Sumber", "Kampanye", "URL Artikel", "URL Gambar", "Catatan Review",
]

PRICE_CUE_RE = re.compile(r"(?i)\b(?:rp\.?|idr|harga|price|promo|normal|asli|reguler|diskon|discount|hemat|off|potongan)\b")
LABEL_NORMAL_RE = re.compile(r"(?i)\b(?:harga\s*(?:normal|asli|reguler|sebelum\s*promo)|normal\s*price|regular\s*price)\b")
LABEL_PROMO_RE = re.compile(r"(?i)\b(?:harga\s*(?:promo|diskon|spesial)|promo\s*price|sale\s*price|price\s*after\s*discount)\b")
DISCOUNT_RE = re.compile(r"(?i)(?<!\d)(\d{1,2}(?:[.,]\d{1,2})?)\s*%")
RP_AMOUNT_RE = re.compile(r"(?i)\b(?:rp\.?|idr)\s*(\d[\d.,]{0,12})")
OCR_MISSING_P_R_AMOUNT_RE = re.compile(r"(?i)(?<![\w])R(?!p)\s*(\d{1,3}(?:[.,:]\d{3})+)(?![\w])")
GROUPED_AMOUNT_RE = re.compile(r"(?<![\w])(\d{1,3}(?:[.,:]\d{3})+)(?![\w])")
PLAIN_AMOUNT_RE = re.compile(r"(?<![\w])([1-9]\d{3,7})(?![\w])")
UNIT_AFTER_RE = re.compile(r"(?i)^\s*(?:ml|g|gr|gram|kg|mg|l|liter|litre|pcs?|buah|lembar|sheets?|pack|sachet|tablet|capsule)\b")


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _read_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    with temp.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    temp.replace(path)


def _safe_image_url(value: str) -> bool:
    """Accept public HTTP(S) image URLs only; never fetch literal private IPs."""
    try:
        p = urlparse(str(value or "").strip())
        host = (p.hostname or "").rstrip(".").lower()
        if p.scheme not in {"http", "https"} or not host or p.username or p.password:
            return False
        if host == "localhost" or host.endswith((".local", ".internal", ".localhost")):
            return False
        try:
            ip = ipaddress.ip_address(host)
            if not ip.is_global:
                return False
        except ValueError:
            # Hostname: only deny local/private suffixes. Public DNS resolution is
            # left to requests; this avoids blocking legitimate public image hosts.
            pass
        return True
    except ValueError:
        return False


def _canonical_article_url(value: object) -> str:
    """Normalize article URLs consistently for OCR-to-product association."""
    try:
        parsed = urlparse(str(value or "").strip())
        if parsed.scheme not in {"http", "https"} or not parsed.hostname:
            return ""
        host = parsed.hostname.lower().removeprefix("www.")
        path = parsed.path.rstrip("/") or "/"
        return urlunparse((parsed.scheme.lower(), host, path, "", parsed.query, ""))
    except ValueError:
        return ""


def _normalize_box(box) -> tuple[float, float, float, float]:
    """Normalize polygon boxes (RapidOCR) and xyxy boxes (PaddleOCR) to xyxy."""
    try:
        points = list(box)
        # PaddleOCR 3.x returns [x0, y0, x1, y1]; RapidOCR returns polygons.
        if len(points) == 4 and all(not isinstance(point, (list, tuple)) and getattr(point, "ndim", 0) == 0 for point in points):
            x0, y0, x1, y1 = (float(point) for point in points)
            return min(x0, x1), min(y0, y1), max(x0, x1), max(y0, y1)
        xs = [float(point[0]) for point in points]
        ys = [float(point[1]) for point in points]
        return min(xs), min(ys), max(xs), max(ys)
    except (TypeError, ValueError, IndexError):
        return 0.0, 0.0, 0.0, 0.0


def _as_list(value) -> list:
    """Convert Python/NumPy OCR outputs to a list without ambiguous truth tests."""
    if value is None:
        return []
    try:
        return list(value)
    except (TypeError, ValueError):
        return []


def _confidence(value) -> float:
    try:
        n = float(value)
    except (TypeError, ValueError):
        return 0.0
    if n > 1.0:
        n /= 100.0
    return min(1.0, max(0.0, n))


def create_ocr_engine(preferred: str = "auto"):
    """Create a local OCR engine; prefer PaddleOCR, then RapidOCR/Tesseract fallbacks."""
    preferred = (preferred or "auto").strip().casefold()
    factories = {
        "paddleocr": ("PaddleOCR",),
        "rapidocr": ("RapidOCR",),
        "tesseract": ("Tesseract",),
        "auto": ("PaddleOCR", "RapidOCR", "Tesseract"),
    }
    if preferred not in factories:
        raise ValueError(f"Backend OCR tidak dikenal: {preferred}")
    errors = []
    for name in factories[preferred]:
        try:
            if name == "PaddleOCR":
                from paddleocr import PaddleOCR
                engine = PaddleOCR(
                    lang="en",
                    device="cpu",
                    use_doc_orientation_classify=False,
                    use_doc_unwarping=False,
                    use_textline_orientation=True,
                )
                return "paddleocr", engine
            if name == "RapidOCR":
                from rapidocr import RapidOCR
                return "rapidocr", RapidOCR()
            import pytesseract
            pytesseract.get_tesseract_version()
            return "tesseract", pytesseract
        except Exception as exc:
            errors.append(f"{name}: {type(exc).__name__}: {exc}")
    if preferred != "auto":
        raise RuntimeError(f"Backend OCR {preferred} tidak tersedia. " + " | ".join(errors))
    raise RuntimeError(
        "Mesin OCR lokal belum tersedia. Jalankan Install_Dependency.bat atau "
        "pasang PaddleOCR/PaddlePaddle. Fallback yang dicoba: " + " | ".join(errors)
    )


def _prepare_image(image: Image.Image) -> Image.Image:
    """Upscale and improve local contrast before OCR without losing colour layout."""
    image = ImageOps.exif_transpose(image).convert("RGB")
    largest = max(image.size)
    if largest < 1800:
        scale = min(2.5, 1800 / max(1, largest))
        image = image.resize((max(1, int(image.width * scale)), max(1, int(image.height * scale))), Image.Resampling.LANCZOS)
    image = ImageOps.autocontrast(image, cutoff=1)
    image = ImageEnhance.Contrast(image).enhance(1.12)
    image = ImageEnhance.Sharpness(image).enhance(1.2)
    # OpenCV CLAHE helps low-contrast small prices. It's optional so the Pillow
    # preprocessing above remains usable if cv2 is unavailable.
    try:
        import cv2
        import numpy as np
        rgb = np.asarray(image)
        lab = cv2.cvtColor(rgb, cv2.COLOR_RGB2LAB)
        light, a, b = cv2.split(lab)
        light = cv2.createCLAHE(clipLimit=1.6, tileGridSize=(8, 8)).apply(light)
        adjusted = cv2.cvtColor(cv2.merge((light, a, b)), cv2.COLOR_LAB2RGB)
        image = Image.fromarray(adjusted)
    except Exception:
        pass
    return image


def recognize_image(image_bytes: bytes, engine, backend: str) -> tuple[list[dict], int, int]:
    with Image.open(io.BytesIO(image_bytes)) as opened:
        image = _prepare_image(opened)
    width, height = image.size
    if backend == "paddleocr":
        import numpy as np
        # PaddleOCR expects an OpenCV-style BGR array for ndarray input.
        bgr = np.asarray(image)[:, :, ::-1].copy()
        result_items = engine.predict(bgr)
        lines = []
        for item in result_items or []:
            try:
                texts = item.get("rec_texts", [])
                scores = item.get("rec_scores", [])
                boxes = item.get("rec_boxes", [])
            except Exception:
                try:
                    texts = item["rec_texts"]
                    scores = item["rec_scores"]
                    boxes = item["rec_boxes"]
                except Exception:
                    continue
            texts = _as_list(texts)
            scores = _as_list(scores)
            boxes = _as_list(boxes)
            for idx, text in enumerate(texts):
                if not str(text or "").strip():
                    continue
                box = boxes[idx] if idx < len(boxes) else []
                score = scores[idx] if idx < len(scores) else 0.0
                x0, y0, x1, y1 = _normalize_box(box)
                lines.append({"text": str(text).strip(), "confidence": _confidence(score), "x0": x0, "y0": y0, "x1": x1, "y1": y1})
        return _sort_lines(lines), width, height

    if backend == "rapidocr":
        result = engine(image)
        texts = getattr(result, "txts", None)
        scores = getattr(result, "scores", None)
        boxes = getattr(result, "boxes", None)
        texts = [] if texts is None else list(texts)
        scores = [] if scores is None else list(scores)
        boxes = [] if boxes is None else list(boxes)
        lines = []
        for idx, text in enumerate(texts):
            if not str(text or "").strip():
                continue
            box = boxes[idx] if idx < len(boxes) else []
            score = scores[idx] if idx < len(scores) else 0.0
            x0, y0, x1, y1 = _normalize_box(box)
            lines.append({"text": str(text).strip(), "confidence": _confidence(score), "x0": x0, "y0": y0, "x1": x1, "y1": y1})
        return _sort_lines(lines), width, height

    # Local Tesseract fallback. It is not installed through pip; only used when
    # RapidOCR is absent but the Tesseract executable is already installed.
    from pytesseract import Output
    import pytesseract
    available_langs = set(pytesseract.get_languages(config=""))
    lang = "ind+eng" if {"ind", "eng"}.issubset(available_langs) else ("eng" if "eng" in available_langs else "osd")
    data = pytesseract.image_to_data(image, lang=lang, config="--psm 11", output_type=Output.DICT)
    grouped = {}
    for idx, text in enumerate(data.get("text", [])):
        text = str(text or "").strip()
        if not text:
            continue
        try:
            conf = max(0.0, float(data.get("conf", [0])[idx])) / 100.0
            x0, y0 = float(data["left"][idx]), float(data["top"][idx])
            x1, y1 = x0 + float(data["width"][idx]), y0 + float(data["height"][idx])
        except (TypeError, ValueError, IndexError):
            conf, x0, y0, x1, y1 = 0.0, 0.0, 0.0, 0.0, 0.0
        key = (data.get("block_num", [0])[idx], data.get("par_num", [0])[idx], data.get("line_num", [0])[idx])
        item = grouped.setdefault(key, {"parts": [], "scores": [], "x0": x0, "y0": y0, "x1": x1, "y1": y1})
        item["parts"].append(text)
        item["scores"].append(conf)
        item["x0"] = min(item["x0"], x0); item["y0"] = min(item["y0"], y0)
        item["x1"] = max(item["x1"], x1); item["y1"] = max(item["y1"], y1)
    lines = []
    for item in grouped.values():
        lines.append({"text": " ".join(item["parts"]), "confidence": statistics.mean(item["scores"]) if item["scores"] else 0.0, **{k: item[k] for k in ("x0", "y0", "x1", "y1")}})
    return _sort_lines(lines), width, height


def _sort_lines(lines: list[dict]) -> list[dict]:
    return sorted(lines, key=lambda line: ((float(line.get("y0", 0)) + float(line.get("y1", 0))) / 2, float(line.get("x0", 0))))


def _integer_amount(raw: str) -> int | None:
    return parse_idr_price(raw)


def extract_prices(text: str) -> list[dict]:
    """Extract plausible IDR amounts, preferring explicit Rp/IDR and grouped numbers."""
    text = str(text or "")
    matches = []
    occupied = []
    for match in RP_AMOUNT_RE.finditer(text):
        price = _integer_amount(match.group(1))
        if price is not None and 100 <= price <= 100_000_000:
            matches.append({"price": price, "start": match.start(), "end": match.end(), "raw": match.group(0), "explicit_currency": True})
            occupied.append((match.start(), match.end()))
    # OCR sometimes drops the P in Rp or reads the thousands separator as ':'.
    # Treat the single-letter R form as a candidate only (it is never auto-approved).
    for match in OCR_MISSING_P_R_AMOUNT_RE.finditer(text):
        if any(match.start() < end and match.end() > start for start, end in occupied):
            continue
        price = _integer_amount(match.group(1))
        if price is not None and 1_000 <= price <= 100_000_000:
            matches.append({"price": price, "start": match.start(), "end": match.end(), "raw": match.group(0), "explicit_currency": False})
            occupied.append((match.start(), match.end()))
    for match in GROUPED_AMOUNT_RE.finditer(text):
        if any(match.start() < end and match.end() > start for start, end in occupied):
            continue
        price = _integer_amount(match.group(1))
        if price is not None and 1_000 <= price <= 100_000_000:
            matches.append({"price": price, "start": match.start(), "end": match.end(), "raw": match.group(1), "explicit_currency": False})
    # Unformatted numbers are accepted only when a nearby currency/price cue exists.
    for match in PLAIN_AMOUNT_RE.finditer(text):
        if any(match.start() < end and match.end() > start for start, end in occupied):
            continue
        value = int(match.group(1))
        if 1_900 <= value <= 2_099:  # avoid campaign years mistaken for prices
            continue
        nearby = text[max(0, match.start() - 30):min(len(text), match.end() + 18)]
        if not re.search(r"(?i)\b(?:rp\.?|idr|harga|price|promo|normal|asli|reguler)\b", nearby):
            continue
        after = text[match.end():match.end() + 16]
        if UNIT_AFTER_RE.search(after):
            continue
        matches.append({"price": value, "start": match.start(), "end": match.end(), "raw": match.group(1), "explicit_currency": False})
    matches.sort(key=lambda item: item["start"])
    # Preserve repeated amounts at distinct locations, but suppress duplicate regex matches.
    out, seen_pos = [], set()
    for item in matches:
        pos = (item["start"], item["end"])
        if pos in seen_pos:
            continue
        seen_pos.add(pos); out.append(item)
    return out


def extract_discount_claim(text: str) -> float | None:
    match = DISCOUNT_RE.search(str(text or ""))
    if not match:
        return None
    try:
        value = float(match.group(1).replace(",", "."))
    except ValueError:
        return None
    return value if 0 < value <= 100 else None


def _labelled_price(text: str, label_re: re.Pattern) -> int | None:
    label = label_re.search(text)
    if not label:
        return None
    after = text[label.end():label.end() + 80]
    prices = extract_prices(after)
    return prices[0]["price"] if prices else None


def interpret_price_pair(text: str, *, allow_unlabelled_pair: bool = False) -> dict:
    """Interpret price roles; spatial OCR can suggest unlabeled pairs for human review only."""
    prices = extract_prices(text)
    unique_prices = list(dict.fromkeys(item["price"] for item in prices))
    normal = _labelled_price(text, LABEL_NORMAL_RE)
    promo = _labelled_price(text, LABEL_PROMO_RE)
    claim = extract_discount_claim(text)
    calculated = None
    note = "Harga belum dapat ditentukan dari teks OCR; perlu dibaca pada poster."

    def plausible_pair(high: int, low: int) -> tuple[float | None, bool]:
        pct = round((high - low) / high * 100, 1) if high else None
        if pct is None or not (2.0 <= pct <= 80.0):
            return pct, False
        if claim is not None and abs(pct - claim) > 3.0:
            return pct, False
        return pct, True

    if normal and promo:
        if normal > promo:
            calculated = round((normal - promo) / normal * 100, 1)
            if claim is None or abs(calculated - claim) <= 3.0:
                note = "Harga normal dan promo dikenali dari label OCR; tetap cocokkan dengan poster."
            else:
                note = "Label harga terbaca, tetapi hitungan tidak cocok dengan klaim diskon; periksa poster."
        else:
            normal = promo = None
            note = "Label/urutan harga konflik; periksa gambar sebelum menyetujui."
    elif len(unique_prices) == 2 and normal in unique_prices and promo is None:
        other = next((price for price in unique_prices if price != normal), None)
        if other is not None and other < normal:
            calculated, ok = plausible_pair(normal, other)
            if ok:
                promo = other
                note = "Harga normal berlabel; harga promo kandidat dari angka lebih rendah di tile yang sama. WAJIB cocokkan poster."
            else:
                calculated = None
                note = "Harga normal berlabel tetapi harga lain tidak cocok dengan klaim/aturan diskon; periksa poster."
        else:
            note = "Harga normal berlabel terdeteksi; harga promo belum jelas."
    elif len(unique_prices) == 2 and promo in unique_prices and normal is None:
        other = next((price for price in unique_prices if price != promo), None)
        if other is not None and other > promo:
            calculated, ok = plausible_pair(other, promo)
            if ok:
                normal = other
                note = "Harga promo berlabel; harga normal kandidat dari angka lebih tinggi di tile yang sama. WAJIB cocokkan poster."
            else:
                calculated = None
                note = "Harga promo berlabel tetapi harga lain tidak cocok dengan klaim/aturan diskon; periksa poster."
        else:
            note = "Harga promo berlabel terdeteksi; harga normal belum jelas."
    elif len(unique_prices) == 2:
        high, low = max(unique_prices), min(unique_prices)
        calculated, ok = plausible_pair(high, low)
        if claim is not None and ok:
            normal, promo = high, low
            note = "Pasangan harga lokal cocok dengan klaim diskon; tetap konfirmasi visual."
        elif claim is not None:
            calculated = None
            note = "Dua harga OCR tidak cocok dengan klaim diskon; tidak dipasangkan otomatis."
        elif allow_unlabelled_pair and ok:
            normal, promo = high, low
            note = "Kandidat lokal: angka terbesar dianggap harga normal dan terkecil harga promo; WAJIB cocokkan poster sebelum TERIMA."
        else:
            calculated = None
            note = "Dua harga terdeteksi tanpa label yang cukup; pasangan dibiarkan kosong."
    elif normal:
        note = "Harga normal berlabel terdeteksi; harga promo belum jelas."
    elif promo:
        note = "Harga promo berlabel terdeteksi; harga normal belum jelas."
    elif len(unique_prices) > 1:
        note = "Lebih dari dua harga dalam konteks; area mungkin mencakup SKU lain. Pasangan dibiarkan kosong."

    if normal and promo and normal > promo:
        calculated = round((normal - promo) / normal * 100, 1)
    elif normal and promo and promo >= normal:
        normal = promo = None
        calculated = None
        note = "Label/urutan harga konflik; periksa gambar sebelum menyetujui."
    return {
        "prices": unique_prices,
        "normal_price_candidate": normal or "",
        "promo_price_candidate": promo or "",
        "discount_claim_percent": claim if claim is not None else "",
        "discount_calculated_percent": calculated if calculated is not None else "",
        "candidate_note": note,
    }


def _line_center(line: dict) -> tuple[float, float]:
    return ((float(line.get("x0", 0)) + float(line.get("x1", 0))) / 2, (float(line.get("y0", 0)) + float(line.get("y1", 0))) / 2)


def make_price_windows(lines: list[dict], width: int, height: int) -> list[dict]:
    windows = []
    seen = set()
    if not lines or width <= 0 or height <= 0:
        return windows
    sorted_lines = _sort_lines(lines)
    for i, line in enumerate(sorted_lines):
        this_text = str(line.get("text") or "").strip()
        if not this_text or not extract_prices(this_text):
            continue
        cx, cy = _line_center(line)
        neighborhood = []
        for other in sorted_lines:
            ox, oy = _line_center(other)
            dx, dy = abs(ox - cx) / width, abs(oy - cy) / height
            # Full catalog pages contain many SKUs in a grid. The older wide
            # window (19% width x 15% height) mixed neighbouring products and
            # left almost every price pair blank. A tighter local tile window
            # reduces cross-product contamination before a pair is proposed.
            if dx <= 0.12 and dy <= 0.075:
                neighborhood.append(other)
        neighborhood = sorted(neighborhood, key=lambda x: (_line_center(x)[1], _line_center(x)[0]))
        context = " | ".join(str(x.get("text") or "").strip() for x in neighborhood if str(x.get("text") or "").strip())
        # Pairing here is only a suggestion because the spatial window is local;
        # Excel review remains mandatory before prices affect the catalog.
        parsed = interpret_price_pair(context, allow_unlabelled_pair=True)
        if not parsed["prices"]:
            continue
        line_conf = statistics.mean(_confidence(x.get("confidence")) for x in neighborhood) if neighborhood else _confidence(line.get("confidence"))
        # Quantized spatial position makes ids stable while keeping adjacent products separate.
        bucket_x = int(cx / max(1, width * 0.10))
        bucket_y = int(cy / max(1, height * 0.055))
        signature = json.dumps([context.casefold(), bucket_x, bucket_y, parsed["prices"]], ensure_ascii=False, sort_keys=True)
        if signature in seen:
            continue
        seen.add(signature)
        windows.append({
            "context": context,
            "center_x": cx,
            "center_y": cy,
            "ocr_confidence": round(line_conf, 4),
            "price_values": parsed["prices"],
            **{k: parsed[k] for k in ("normal_price_candidate", "promo_price_candidate", "discount_claim_percent", "discount_calculated_percent", "candidate_note")},
            "lines_in_context": neighborhood,
            "bucket_x": bucket_x,
            "bucket_y": bucket_y,
        })
    return windows


def _matching_score(product_name: str, context: str) -> tuple[float, str]:
    name_key = identity_key(product_name)
    text_key = identity_key(context)
    if not name_key or not text_key:
        return 0.0, "Nama tidak terbaca cukup jelas."
    name_tokens = set(name_key.split())
    text_tokens = set(text_key.split())
    coverage = len(name_tokens & text_tokens) / max(1, len(name_tokens))
    similarity = product_name_similarity(product_name, context)
    score = round(0.55 * similarity + 0.45 * coverage, 4)
    expected_sizes = set(product_name_size_signature(product_name))
    seen_sizes = set(product_name_size_signature(context))
    if expected_sizes and not (expected_sizes & seen_sizes):
        score *= 0.72
        note = "Ukuran SKU tidak terbaca/cocok di konteks OCR; perlu konfirmasi."
    else:
        note = "Ukuran cocok atau tidak dicantumkan pada nama katalog."
    return round(score, 4), note


def _suggest_product(context: str, page_products: list[dict]) -> tuple[str, float, str]:
    best_name, best_score, best_note = "", 0.0, "Tidak ada produk artikel yang cocok kuat; isi nama produk setelah melihat poster."
    for product in page_products:
        if str(product.get("product_validation_status") or "eligible").casefold() not in {"eligible", "", "valid"}:
            continue
        name = str(product.get("product_text") or "").strip()
        score, note = _matching_score(name, context)
        if score > best_score:
            best_name, best_score, best_note = name, score, note
    if best_score < 0.48:
        return "", best_score, "Nama produk belum cocok kuat; isi SKU setelah membuka poster asli."
    return best_name, best_score, best_note


def build_candidates(
    image_record: dict,
    lines: list[dict],
    width: int,
    height: int,
    page_products: list[dict],
    fallback_products: list[dict] | None = None,
) -> list[dict]:
    """Create local price candidates; prefer same-article names over catalog suggestions."""
    raw_candidates = []
    for window in make_price_windows(lines, width, height):
        context = window["context"]
        suggested, score, match_note = _suggest_product(context, page_products)
        if not suggested and fallback_products:
            fallback_name, fallback_score, fallback_note = _suggest_product(context, fallback_products)
            if fallback_name and fallback_score >= 0.55:
                suggested, score = fallback_name, fallback_score
                match_note = "Saran nama dari katalog lain, bukan bukti langsung poster; verifikasi nama/ukuran. " + fallback_note
        price_values = window["price_values"]
        seed = "|".join([
            str(image_record.get("image_url") or ""),
            str(window.get("bucket_x")), str(window.get("bucket_y")),
            identity_key(context), ",".join(map(str, price_values)),
        ])
        candidate_id = hashlib.sha256(seed.encode("utf-8")).hexdigest()[:14].upper()
        raw_candidates.append({
            "candidate_id": candidate_id,
            "source_name": image_record.get("source_name", ""),
            "campaign_title": image_record.get("campaign_title", ""),
            "article_url": image_record.get("article_url", ""),
            "image_url": image_record.get("image_url", ""),
            "image_alt": image_record.get("image_alt", ""),
            "ocr_snippet": context,
            "price_values_json": json.dumps(price_values, ensure_ascii=False),
            "normal_price_candidate": window["normal_price_candidate"],
            "promo_price_candidate": window["promo_price_candidate"],
            "discount_claim_percent": window["discount_claim_percent"],
            "discount_calculated_percent": window["discount_calculated_percent"],
            "ocr_confidence": window["ocr_confidence"],
            "matched_product_name": suggested,
            "matched_product_score": score,
            "match_note": match_note,
            "candidate_note": window["candidate_note"],
            "_bucket_x": int(window.get("bucket_x", 0)),
            "_bucket_y": int(window.get("bucket_y", 0)),
            "_context_lines": {str(line.get("text") or "").strip() for line in window.get("lines_in_context", []) if str(line.get("text") or "").strip()},
        })

    # OCR produces one window around each price line. Adjacent windows often
    # describe the same offer, e.g. one sees normal+promo and the next sees
    # normal+promo+discount. Merge only spatially adjacent windows with the same
    # price pair; if two distinct product names are confidently recognized, keep
    # them separate to avoid merging two separate SKUs sold at the same price.
    merged = []
    for candidate in raw_candidates:
        values = (candidate.get("price_values_json"), candidate.get("normal_price_candidate"), candidate.get("promo_price_candidate"))
        claim = candidate.get("discount_claim_percent")
        product_key = identity_key(candidate.get("matched_product_name"))
        match_index = None
        for idx, previous in enumerate(merged):
            previous_values = (previous.get("price_values_json"), previous.get("normal_price_candidate"), previous.get("promo_price_candidate"))
            if candidate.get("image_url") != previous.get("image_url") or values != previous_values:
                continue
            if abs(candidate["_bucket_x"] - previous["_bucket_x"]) > 1 or abs(candidate["_bucket_y"] - previous["_bucket_y"]) > 1:
                continue
            previous_product_key = identity_key(previous.get("matched_product_name"))
            if product_key and previous_product_key and product_key != previous_product_key:
                continue
            previous_claim = previous.get("discount_claim_percent")
            if claim not in (None, "") and previous_claim not in (None, "") and str(claim) != str(previous_claim):
                continue
            match_index = idx
            break
        if match_index is None:
            merged.append(candidate)
            continue
        previous = merged[match_index]
        # Keep the strongest product match and OCR confidence, but retain the
        # richer context/discount clue from the other adjacent window.
        if (float(candidate.get("matched_product_score") or 0), float(candidate.get("ocr_confidence") or 0)) > (float(previous.get("matched_product_score") or 0), float(previous.get("ocr_confidence") or 0)):
            chosen, other = candidate, previous
        else:
            chosen, other = previous, candidate
        chosen = dict(chosen)
        combined_lines = chosen.get("_context_lines", set()) | other.get("_context_lines", set())
        chosen["ocr_snippet"] = " | ".join(sorted(combined_lines)) or chosen.get("ocr_snippet", "")
        chosen["_context_lines"] = combined_lines
        if chosen.get("discount_claim_percent") in (None, "") and other.get("discount_claim_percent") not in (None, ""):
            chosen["discount_claim_percent"] = other["discount_claim_percent"]
            chosen["discount_calculated_percent"] = other.get("discount_calculated_percent")
        if chosen.get("candidate_note") and other.get("candidate_note") and chosen.get("candidate_note") != other.get("candidate_note"):
            chosen["candidate_note"] = chosen.get("candidate_note") + " Window OCR bertumpang tindih digabung; tetap periksa poster."
        merged[match_index] = chosen

    for candidate in merged:
        candidate.pop("_bucket_x", None)
        candidate.pop("_bucket_y", None)
        candidate.pop("_context_lines", None)
    return merged


def _download_image(session: requests.Session, image_url: str) -> tuple[bytes, str]:
    if not _safe_image_url(image_url):
        raise ValueError("URL gambar tidak aman/tidak valid")
    with session.get(image_url, timeout=REQUEST_TIMEOUT, stream=True, headers={"User-Agent": USER_AGENT}) as response:
        response.raise_for_status()
        ctype = (response.headers.get("Content-Type") or "").split(";")[0].strip().lower()
        if ctype and not ctype.startswith("image/"):
            raise ValueError(f"URL tidak mengembalikan konten gambar ({ctype})")
        chunks = []
        total = 0
        for chunk in response.iter_content(chunk_size=64 * 1024):
            if not chunk:
                continue
            total += len(chunk)
            if total > MAX_IMAGE_BYTES:
                raise ValueError("Gambar melebihi batas 20 MB")
            chunks.append(chunk)
    body = b"".join(chunks)
    if not body:
        raise ValueError("Gambar kosong")
    return body, hashlib.sha256(body).hexdigest()


def _line_records(cache_row: dict) -> list[dict]:
    try:
        lines = json.loads(cache_row.get("lines_json") or "[]")
        return lines if isinstance(lines, list) else []
    except json.JSONDecodeError:
        return []


def _load_old_review(path: Path) -> dict[str, dict]:
    if not path.exists():
        return {}
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if "Kandidat OCR" not in wb.sheetnames:
            wb.close(); return {}
        ws = wb["Kandidat OCR"]
        headers = [str(c or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        output = {}
        for values in ws.iter_rows(min_row=2, values_only=True):
            row = dict(zip(headers, values))
            cid = str(row.get("ID Kandidat") or "").strip()
            if cid:
                output[cid] = row
        wb.close()
        return output
    except Exception:
        return {}


def make_review_workbook(candidates: list[dict], path: Path = REVIEW_XLSX, prior: dict[str, dict] | None = None) -> None:
    prior = prior or {}
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Kandidat OCR"
    ws.append(REVIEW_HEADERS)
    for candidate in candidates:
        cid = candidate["candidate_id"]
        old = prior.get(cid, {})
        # User-editable decision/values take priority when the same source image is reprocessed.
        suggested = candidate.get("matched_product_name") or ""
        decision = str(old.get("Keputusan") or "CEK").strip().upper()
        if decision not in {"CEK", "TERIMA", "TOLAK"}:
            decision = "CEK"
        product_approved = old.get("Nama Produk Disetujui") or suggested
        normal_approved = old.get("Harga Normal Disetujui (Rp)")
        if normal_approved in (None, ""):
            normal_approved = candidate.get("normal_price_candidate") or ""
        promo_approved = old.get("Harga Promo Disetujui (Rp)")
        if promo_approved in (None, ""):
            promo_approved = candidate.get("promo_price_candidate") or ""
        row_values = [
            cid, decision, product_approved, suggested,
            normal_approved, promo_approved,
            candidate.get("normal_price_candidate") or "", candidate.get("promo_price_candidate") or "",
            candidate.get("discount_claim_percent") or "", candidate.get("discount_calculated_percent") or "",
            candidate.get("match_note", ""), candidate.get("ocr_confidence", ""), candidate.get("ocr_snippet", ""),
            candidate.get("source_name", ""), candidate.get("campaign_title", ""), candidate.get("article_url", ""),
            candidate.get("image_url", ""), old.get("Catatan Review") or candidate.get("candidate_note", ""),
        ]
        ws.append(row_values)
        # Make both provenance URLs clickable in Excel so reviewers can open the
        # original poster/article directly instead of copying long URLs manually.
        current_row = ws.max_row
        for column in (16, 17):
            cell = ws.cell(current_row, column)
            url = str(cell.value or "").strip()
            if url.startswith(("http://", "https://")):
                cell.hyperlink = url
                cell.style = "Hyperlink"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 42
    widths = [15, 13, 38, 38, 22, 22, 20, 20, 16, 16, 45, 14, 76, 22, 48, 60, 70, 68]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for idx in (5, 6, 7, 8):
            row[idx - 1].number_format = '#,##0'
        for idx in (9, 10):
            row[idx - 1].number_format = '0.0"%"'
        # Editable columns are highlighted using a pale yellow fill.
        for idx in (2, 3, 5, 6, 18):
            row[idx - 1].fill = PatternFill("solid", fgColor="FFF2CC")
    decision_validation = DataValidation(type="list", formula1='"CEK,TERIMA,TOLAK"', allow_blank=False)
    decision_validation.error = "Pilih CEK, TERIMA, atau TOLAK."
    decision_validation.errorTitle = "Keputusan OCR tidak valid"
    ws.add_data_validation(decision_validation)
    if ws.max_row >= 2:
        decision_validation.add(f"B2:B{ws.max_row}")

    helpws = wb.create_sheet("Petunjuk")
    tips = [
        ["REVIEW HARGA DARI POSTER PROMO (OCR LOKAL)"],
        ["1. Buka URL Gambar untuk melihat poster asli; OCR dapat salah membaca digit atau menggabungkan dua SKU."],
        ["2. Periksa Nama Produk Disetujui dan Harga Normal/Promo Disetujui. Kolom berwarna kuning adalah kolom yang boleh diedit."],
        ["3. Pilih TERIMA hanya bila SKU, ukuran/varian, dan peran harga sudah cocok secara visual. Pilih TOLAK bila kandidat salah; selain itu biarkan CEK."],
        ["4. Jika harga normal dicoret dan harga promo tidak berlabel, pastikan konteks poster/klaim diskon mendukung pasangan tersebut sebelum TERIMA."],
        ["5. Setelah menyimpan dan menutup Excel, jalankan: python engine\\promo_poster_ocr.py --import-reviewed-only"],
        ["6. Lalu perbarui katalog: python engine\\enrich_existing_promo_catalog_multisource.py. Harga dari poster hanya dipakai setelah TERIMA."],
        ["7. Tidak ada AI generatif/tebakan harga. OCR menyalin teks/angka; keputusan harga tetap membutuhkan konteks bukti poster."],
        ["8. Kandidat yang tidak punya kecocokan produk tetap bisa ditinjau, tetapi isi Nama Produk Disetujui sesuai SKU katalog dan cocokkan URL artikelnya."],
    ]
    for tip in tips:
        helpws.append(tip)
    helpws.column_dimensions["A"].width = 150
    for row in helpws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row == 1:
                cell.font = Font(bold=True, size=13, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="244062")
    wb.save(path)


def _load_image_index(path: Path = IMAGE_INDEX) -> list[dict]:
    rows = _read_csv(path)
    deduped = {}
    for row in rows:
        url = str(row.get("image_url") or "").strip()
        if not url or not _safe_image_url(url):
            continue
        if url not in deduped:
            deduped[url] = row
        else:
            # Keep richer campaign metadata if duplicated image URL appears under another article.
            current = deduped[url]
            if not current.get("article_url") and row.get("article_url"):
                deduped[url] = row
    return list(deduped.values())


def _load_page_products(path: Path = PRODUCTS_CSV) -> dict[str, list[dict]]:
    by_url = defaultdict(list)
    for row in _read_csv(path):
        urls = {row.get("source_url"), row.get("source_page_url")}
        for url in urls:
            key = _canonical_article_url(url)
            if key:
                by_url[key].append(row)
    return by_url


def _load_catalog_products(path: Path = PRODUCTS_CSV) -> list[dict]:
    """Load eligible SKU names only; their prices are never treated as OCR evidence."""
    output = []
    for row in _read_csv(path):
        status = str(row.get("product_validation_status") or "eligible").casefold()
        product = str(row.get("product_text") or "").strip()
        if status not in {"eligible", "", "valid"} or not product:
            continue
        output.append({
            "product_text": product,
            "product_validation_status": status,
            "merchant_name": str(row.get("merchant_name") or ""),
            "source_name": str(row.get("source_name") or ""),
        })
    seen, unique = set(), []
    for row in output:
        key = identity_key(row["product_text"])
        if key and key not in seen:
            seen.add(key)
            unique.append(row)
    return unique


def _catalog_fallback_for_campaign(catalog: list[dict], campaign_title: str) -> list[dict]:
    """Return possible names from the retailer named in the campaign title."""
    campaign_key = identity_key(campaign_title)
    if not campaign_key:
        return []
    campaign_tokens = set(campaign_key.split())
    merchant_keys = {
        identity_key(row.get("merchant_name"))
        for row in catalog
        if identity_key(row.get("merchant_name"))
    }
    matched = {key for key in merchant_keys if key in campaign_tokens}
    if not matched:
        return []
    return [row for row in catalog if identity_key(row.get("merchant_name")) in matched]


def run_ocr(limit: int = 0, force: bool = False, backend_preference: str = "auto") -> dict:
    images = _load_image_index()
    if not images:
        raise FileNotFoundError(f"Daftar gambar poster tidak ditemukan/kosong: {IMAGE_INDEX}. Jalankan collector terlebih dahulu.")
    old_cache = {str(r.get("image_url") or ""): r for r in _read_csv(CACHE_CSV)}
    if force:
        to_process = images[:limit] if limit > 0 else images
    else:
        # Retry transient download/OCR failures next run. Successful results and
        # no-text images remain cached to avoid repeatedly processing the same poster.
        to_process = [
            r for r in images
            if str(r.get("image_url") or "") not in old_cache
            or str(old_cache.get(str(r.get("image_url") or ""), {}).get("ocr_status") or "").casefold() == "failed"
        ]
        if limit > 0:
            to_process = to_process[:limit]
    if not to_process:
        print("OCR: semua gambar pada indeks sudah memiliki cache; gunakan --force untuk memproses ulang.")
        # Even with no new images, the user may have edited the review workbook.
        # Import those decisions so the following enricher run sees the approval.
        if REVIEW_XLSX.exists():
            import_stats = import_reviewed_prices(REVIEW_XLSX, OVERRIDES_CSV)
            print(f"Koreksi poster disetujui/ditambahkan: {import_stats.get('accepted_or_updated', 0)}")
            print(f"Koreksi poster dinonaktifkan: {import_stats.get('deactivated', 0)}")
            print(f"Koreksi poster aktif: {import_stats.get('active_overrides', 0)}")
        else:
            import_stats = {"active_overrides": 0}
        return {
            "images_indexed": len(images), "images_processed": 0,
            "images_cached": len([r for r in images if r.get("image_url") in old_cache]),
            "ocr_ok": 0, "ocr_failed": 0, "candidate_count": 0,
            "approved_override_count": import_stats.get("active_overrides", 0),
        }

    backend, engine = create_ocr_engine(backend_preference)
    session = requests.Session()
    processed = ok = failed = 0
    try:
        for idx, record in enumerate(to_process, 1):
            url = str(record.get("image_url") or "")
            previous_cache = old_cache.get(url)
            cache_to_store = None
            cache_row = dict(record)
            cache_row.update({"image_sha256": "", "image_width": "", "image_height": "", "ocr_status": "failed", "ocr_backend": backend, "ocr_error": "", "processed_at": now_iso(), "lines_json": "[]"})
            try:
                body, image_hash = _download_image(session, url)
                lines, width, height = recognize_image(body, engine, backend)
                cache_row.update({
                    "image_sha256": image_hash, "image_width": width, "image_height": height,
                    "ocr_status": "ok" if lines else "no_text", "ocr_backend": backend,
                    "processed_at": now_iso(), "lines_json": json.dumps(lines, ensure_ascii=False, separators=(",", ":")),
                })
                if lines:
                    ok += 1
                else:
                    failed += 1
            except Exception as exc:
                cache_row["ocr_error"] = f"{type(exc).__name__}: {exc}"[:500]
                cache_row["ocr_backend"] = backend
                cache_row["processed_at"] = now_iso()
                failed += 1
                # Do not destroy a previous successful OCR result if a forced
                # refresh temporarily fails (e.g. the poster host is down).
                if previous_cache and previous_cache.get("ocr_status") == "ok" and _line_records(previous_cache):
                    cache_to_store = dict(previous_cache)
                    cache_to_store["ocr_error"] = ("Refresh gagal; hasil OCR lama dipertahankan. " + cache_row["ocr_error"])[:500]
            if cache_to_store is None:
                cache_to_store = cache_row
            old_cache[url] = {field: str(cache_to_store.get(field) or "") for field in CACHE_FIELDS}
            processed += 1
            if idx == 1 or idx % 10 == 0 or idx == len(to_process):
                print(f"OCR gambar {idx}/{len(to_process)} | teks terbaca={ok} | gagal/kosong={failed}", flush=True)
            time.sleep(0.05)
    finally:
        session.close()

    # Write cache only for current images and previous cached entries; retaining prior
    # rows allows OCR to be incremental across runs even when an article disappears.
    _write_csv(CACHE_CSV, [old_cache[k] for k in sorted(old_cache)], CACHE_FIELDS)
    cache_rows = [old_cache.get(str(r.get("image_url") or ""), {}) for r in images]
    line_rows = []
    candidates = []
    page_products = _load_page_products()
    catalog_products = _load_catalog_products()
    for image_record, cache_row in zip(images, cache_rows):
        lines = _line_records(cache_row)
        width = int(cache_row.get("image_width") or 0)
        height = int(cache_row.get("image_height") or 0)
        stamp = cache_row.get("processed_at") or ""
        for n, line in enumerate(_sort_lines(lines), 1):
            line_rows.append({
                "source_name": image_record.get("source_name", ""), "campaign_title": image_record.get("campaign_title", ""),
                "article_url": image_record.get("article_url", ""), "image_url": image_record.get("image_url", ""),
                "image_alt": image_record.get("image_alt", ""), "period_status": image_record.get("period_status", ""),
                "line_no": n, "ocr_text": line.get("text", ""), "ocr_confidence": line.get("confidence", ""),
                "x0": line.get("x0", ""), "y0": line.get("y0", ""), "x1": line.get("x1", ""), "y1": line.get("y1", ""),
                "image_width": width, "image_height": height, "processed_at": stamp,
            })
        product_rows = page_products.get(_canonical_article_url(image_record.get("article_url")), [])
        fallback_rows = _catalog_fallback_for_campaign(catalog_products, str(image_record.get("campaign_title") or ""))
        candidates.extend(build_candidates(image_record, lines, width, height, product_rows, fallback_rows))

    _write_csv(LINES_CSV, line_rows, LINE_FIELDS)
    # De-duplicate candidates when one image has repeated identical price windows.
    candidate_map = {}
    for candidate in candidates:
        candidate_map[candidate["candidate_id"]] = candidate
    candidates = sorted(candidate_map.values(), key=lambda row: (row.get("source_name", ""), row.get("campaign_title", ""), row.get("candidate_id", "")))
    _write_csv(CANDIDATES_CSV, candidates, CANDIDATE_FIELDS)
    prior = _load_old_review(REVIEW_XLSX)
    make_review_workbook(candidates, REVIEW_XLSX, prior)
    import_stats = import_reviewed_prices(REVIEW_XLSX, OVERRIDES_CSV)
    return {
        "images_indexed": len(images), "images_processed": processed, "images_cached": len(images) - processed,
        "ocr_ok": ok, "ocr_failed_or_empty": failed, "ocr_backend": backend,
        "line_count": len(line_rows), "candidate_count": len(candidates),
        "candidate_price_pair_count": sum(bool(r.get("normal_price_candidate")) and bool(r.get("promo_price_candidate")) for r in candidates),
        "candidate_product_suggestion_count": sum(bool(r.get("matched_product_name")) for r in candidates),
        "approved_override_count": import_stats.get("active_overrides", 0),
        "cache_csv": str(CACHE_CSV), "lines_csv": str(LINES_CSV), "candidate_csv": str(CANDIDATES_CSV),
        "review_workbook": str(REVIEW_XLSX),
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="OCR poster katalog promo secara lokal; tidak melakukan estimasi harga AI.")
    parser.add_argument("--limit", type=int, default=0, help="Maksimum gambar baru diproses (0=semua yang belum ada cache).")
    parser.add_argument("--force", action="store_true", help="Paksa OCR ulang gambar yang sudah ada cache.")
    parser.add_argument("--backend", choices=("auto", "paddleocr", "rapidocr", "tesseract"), default="auto", help="Mesin OCR lokal (default: PaddleOCR lalu fallback RapidOCR/Tesseract).")
    parser.add_argument("--import-reviewed-only", action="store_true", help="Impor keputusan TERIMA/TOLAK dari workbook tanpa OCR ulang.")
    args = parser.parse_args(argv)
    if args.limit < 0:
        parser.error("--limit harus 0 atau lebih")
    if args.import_reviewed_only:
        stats = import_reviewed_prices(REVIEW_XLSX, OVERRIDES_CSV)
        print(f"Baris review dibaca: {stats['review_rows']}")
        print(f"Koreksi disetujui/ditambahkan: {stats['accepted_or_updated']}")
        print(f"Koreksi dinonaktifkan: {stats['deactivated']}")
        print(f"Baris TERIMA tidak valid: {stats['rejected_invalid']}")
        print(f"Koreksi aktif: {stats['active_overrides']}")
        print(f"Database koreksi: {OVERRIDES_CSV}")
        for err in stats.get("errors", [])[:20]:
            print(" - " + err)
        return 0 if not stats["rejected_invalid"] else 2
    try:
        stats = run_ocr(limit=args.limit, force=args.force, backend_preference=args.backend)
    except (FileNotFoundError, RuntimeError, ValueError) as exc:
        print(f"ERROR OCR: {exc}", file=sys.stderr)
        return 2
    for key, value in stats.items():
        print(f"{key}: {value}")
    print("Harga OCR belum dipakai ke katalog sampai keputusan Excel = TERIMA.")
    print("Backend OCR: auto memilih PaddleOCR lokal, lalu fallback RapidOCR/Tesseract jika tersedia.")
    print("Untuk impor manual setelah edit Excel: python engine\\promo_poster_ocr.py --import-reviewed-only")
    print("Setelah itu: python engine\\enrich_existing_promo_catalog_multisource.py")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
