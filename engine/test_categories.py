from __future__ import annotations

import sys
from pathlib import Path

ENGINE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(ENGINE_DIR))

import promo_categories
import promo_multisource_collector as collector
import promo_catalog_all_merchants as momsmoney
import enrich_existing_promo_catalog_multisource as enhancer


def test_taxonomy_has_29_unique_main_categories_and_separates_general_from_review():
    assert len(promo_categories.CATEGORY_TAXONOMY) == 29
    assert len(promo_categories.CATEGORY_NAMES) == len(set(promo_categories.CATEGORY_NAMES)) == 29
    assert promo_categories.CATEGORY_NAMES[-2:] == [
        "Lainnya / Perlu Review",
        "Promo Umum / Multi-Kategori",
    ]
    assert promo_categories.classify_promo_category({"product_text": "Promo JSM Indomaret 23-25 Oktober 2026"}) == (
        "Promo Umum / Multi-Kategori", "Kampanye Umum"
    )
    assert promo_categories.classify_promo_category({"product_text": "Akhirnya"}) == (
        "Lainnya / Perlu Review", "Produk belum teridentifikasi"
    )


def test_common_product_types_map_to_expected_category_and_subcategory():
    cases = [
        ("Beras premium 5 kg", "Makanan & Bahan Pokok", "Beras"),
        ("Susu UHT Full Cream 1 L", "Minuman & Produk Susu", "Susu/produk susu"),
        ("Indomie Goreng 5 pcs", "Snack & Makanan Kemasan", "Mi instan"),
        ("All Fruit Fuji Apples 1 kg", "Produk Segar & Beku", "Buah"),
        ("Paket Ayam Goreng HokBen", "Restoran & Makanan Siap Saji", "Menu/paket restoran"),
        ("Lifebuoy Sabun Mandi", "Perawatan Tubuh & Mandi", "Sabun mandi/body wash"),
        ("Pantene Shampoo", "Perawatan Rambut & Kulit", "Sampo"),
        ("Wardah Lipstick", "Kecantikan & Kosmetik", "Produk bibir"),
        ("Colgate Optic White 100 g", "Perawatan Gigi & Mulut", "Pasta gigi"),
        ("Rinso deterjen cair", "Cuci Pakaian", "Deterjen"),
        ("Sunlight Jeruk 755 ml", "Kebersihan Rumah Tangga", "Sabun cuci piring"),
        ("Tisu wajah Paseo", "Tisu & Produk Kertas", "Tisu wajah"),
        ("Pampers Pants", "Bayi & Anak", "Popok"),
        ("Vitamin C 1000 mg", "Kesehatan & Farmasi", "Vitamin/alat kesehatan/obat bebas"),
        ("Wadah makanan kedap udara", "Rumah Tangga & Dapur", "Peralatan rumah/dapur"),
        ("Whiskas makanan kucing", "Hewan Peliharaan", "Makanan hewan"),
        ("iPhone 18 Pro", "Elektronik & Gadget", "Smartphone/tablet"),
        ("Nike Air Max Shoes", "Fashion & Aksesori", "Sepatu/sandal"),
        ("Oli motor synthetic", "Otomotif", "Oli/perawatan kendaraan"),
        ("Buku tulis 58 lembar", "Pendidikan, Buku & Alat Tulis", "Buku"),
        ("Voucher Belanja Rp100.000", "Voucher & Gift Card", "Voucher/kupon/gift card"),
        ("Cashback 50% Pakai QRIS BCA", "Pembayaran, Cashback & Loyalitas", "Cashback"),
        ("Voucher ongkir Shopee", "Belanja Online & Marketplace", "Marketplace/ongkir"),
        ("Tiket pesawat Jakarta Bali", "Perjalanan & Transportasi", "Tiket/hotel/transportasi"),
        ("Tiket bioskop XXI", "Hiburan & Rekreasi", "Tiket hiburan/rekreasi"),
        ("Paket data Telkomsel", "Telekomunikasi & Digital", "Pulsa/paket data/internet"),
        ("Jasa cuci sofa", "Layanan & Jasa", "Servis/perawatan/jasa"),
        ("ABC Battery Alkaline AA LR-6 8'S", "Elektronik & Gadget", "Baterai/perangkat listrik"),
        ("Vaseline Healthy Bright Lotion 330 ml", "Perawatan Tubuh & Mandi", "Body lotion/serum"),
    ]
    for product, expected_category, expected_subcategory in cases:
        assert promo_categories.classify_promo_category({"product_text": product}) == (
            expected_category, expected_subcategory
        ), product


def test_product_category_is_prioritized_over_campaign_discount_language():
    row = {
        "product_text": "Susu UHT Full Cream 1 L",
        "offer_text": "Cashback 50% pakai QRIS BCA",
        "campaign_title": "Promo Indomaret Diskon Besar Minggu Ini",
    }
    assert promo_categories.classify_promo_category(row)[0] == "Minuman & Produk Susu"


def test_multisource_and_momsmoney_adapters_emit_category_fields():
    meta = collector.parse_promo_metadata("Promo Supermarket 23-25 Oktober 2026", "", "Berlaku 23-25 Oktober 2026")
    row = collector.make_product_row(
        source_name="Hemat.id",
        product_text="Rinso deterjen cair 700 ml",
        offer_text="Harga Promo Rp15.000",
        title="Promo Indomaret 23-25 Oktober 2026",
        merchant="Indomaret", meta=meta,
        source_url="https://example.test/promo", source_page_url="https://example.test/promo",
        source_item_url="https://example.test/item", raw_line="Rinso deterjen cair 700 ml Harga Promo Rp15.000",
        source_mode="hemat",
    )
    assert row["product_category"] == "Cuci Pakaian"
    assert row["product_subcategory"] == "Deterjen"

    moms = momsmoney.parse_product_line(
        "Susu UHT Full Cream 1 L – Harga Normal Rp20.000, Harga Promo Rp15.000",
        "Promo Harian", "https://example.test/moms-promo"
    )
    assert moms is not None
    assert moms["product_category"] == "Minuman & Produk Susu"
    assert moms["product_subcategory"] == "Susu/produk susu"


def test_clean_excel_rows_include_filterable_category_columns_and_reference_sheets():
    row = {
        "product_text": "Rinso deterjen cair 700 ml",
        "merchant_name": "Indomaret",
        "campaign_title": "Promo Indomaret 23-25 Oktober 2026",
        "offer_text": "Harga Promo Rp15.000",
        "promo_status": "active",
        "period_certainty": "confirmed",
        "valid_from": "2026-10-23",
        "valid_until": "2026-10-25",
        "period_precision": "range",
        "source_url": "https://example.test/promo",
        "verification_status": "aggregator_only",
        "source_name": "Hemat.id",
        "promo_price": 15000,
        "normal_price": None,
    }
    clean = enhancer.make_clean_row(row)
    assert clean["Kategori Produk"] == "Cuci Pakaian"
    assert clean["Subkategori"] == "Deterjen"

    from tempfile import TemporaryDirectory
    from openpyxl import load_workbook

    with TemporaryDirectory() as temp_dir:
        path = Path(temp_dir) / "category_test.xlsx"
        assert enhancer.write_excel_catalog(path, [row], [row], [], [], [], []) is True
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            assert "Ringkasan Kategori" in workbook.sheetnames
            assert "Referensi Kategori" in workbook.sheetnames
            for sheet_name in ("Aktif", "Periode Pasti"):
                headers = [cell.value for cell in workbook[sheet_name][1]]
                assert "Kategori Produk" in headers
                assert "Subkategori" in headers
            ref_sheet = workbook["Referensi Kategori"]
            assert ref_sheet.max_row == 30  # header + 29 categories
        finally:
            workbook.close()


def test_ambiguous_brand_and_flavour_words_do_not_override_product_type():
    cases = {
        "Rejoice Shampoo Hijab 170 ml": ("Perawatan Rambut & Kulit", "Sampo"),
        "Enervon-C Tablet 30s": ("Kesehatan & Farmasi", "Vitamin/alat kesehatan/obat bebas"),
        "Davidoff Cool Water Eau de Toilette 75 ml": ("Kecantikan & Kosmetik", "Parfum/fragrance"),
        "So Klin Eau de Parfum Softergent 700 ml": ("Cuci Pakaian", "Deterjen"),
        "Glade Automatic Spray Cool Water": ("Kebersihan Rumah Tangga", "Pengharum/deodorizer ruangan"),
        "WALL'S Magnum Almond 80 ml": ("Produk Segar & Beku", "Es krim/dessert beku"),
        "My Baby Body Cologne Sweet Floral 100 ml": ("Bayi & Anak", "Perawatan kulit/tubuh bayi"),
        "Sunlight Jeruk Nipis 755 ml": ("Kebersihan Rumah Tangga", "Sabun cuci piring"),
        "Kino Larutan Jeruk 320 ml": ("Minuman & Produk Susu", "Minuman herbal/penyegar"),
    }
    for product, expected in cases.items():
        assert promo_categories.classify_promo_category({"product_text": product}) == expected, product


def test_classifier_never_emits_a_category_outside_the_reference_taxonomy():
    examples = [
        "Beras premium 5 kg",
        "Indomie Goreng 5 pcs",
        "Paket Ayam Goreng HokBen",
        "Promo JSM Indomaret 23-25 Oktober 2026",
        "Akhirnya",
        "Parfum Cool Water",
        "Hanya diskon 50%",
    ]
    for product in examples:
        category, _ = promo_categories.classify_promo_category({"product_text": product})
        assert category in promo_categories.CATEGORY_NAMES, (product, category)


def test_category_classifier_reuses_compiled_regex_patterns():
    cache = promo_categories._COMPILED_PATTERN_CACHE
    cache.clear()

    product = {"product_text": "Rejoice Shampoo Hijab 170 ml"}
    first = promo_categories.classify_promo_category(product)
    compiled_count = len(cache)
    second = promo_categories.classify_promo_category(product)

    assert first == second == ("Perawatan Rambut & Kulit", "Sampo")
    assert compiled_count > 0
    assert len(cache) == compiled_count
