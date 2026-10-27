from __future__ import annotations

import json


def _line(text, x0, y0, x1, y1, confidence=0.96):
    return {
        "text": text,
        "confidence": confidence,
        "x0": x0,
        "y0": y0,
        "x1": x1,
        "y1": y1,
    }


def test_payment_campaign_lines_are_excluded_from_sku_title_but_retained_for_audit():
    from promo_poster_ocr import make_price_windows, classify_promo_context

    lines = [
        _line("NUTRIBRA Pro 900g", 110, 210, 340, 238),
        _line("Aplikasi gopay Bayar Pakai Pay", 110, 240, 360, 268),
        _line("18.400", 120, 350, 210, 370),
        _line("Rp15.900", 120, 375, 235, 397),
    ]
    windows = make_price_windows(lines, 800, 800)
    assert windows
    target = next((w for w in windows if w["price_values"] == [18400, 15900]), None)
    assert target is not None
    safe_name = target["local_product_name_suggestion"].casefold()
    assert "nutribra" in safe_name
    assert "gopay" not in safe_name
    assert target["local_product_name_suggestion_status"] == "DIPISAHKAN DARI TEKS KAMPANYE"
    assert target["local_product_name_suggestion_risk"] == "SEDANG"
    assert target["local_product_name_suggestion_raw"]
    assert target["payment_campaign_cue_count"] >= 1
    promo_context = classify_promo_context(target["context"])
    assert "TEKS PROMO PEMBAYARAN/KAMPANYE" in promo_context["promo_context_tags"]


def test_payment_only_title_region_is_quarantined_not_used_as_sku():
    from promo_poster_ocr import make_price_windows, build_candidates

    lines = [
        _line("Aplikasi gopay Bayar Pakai Pay", 110, 240, 360, 268),
        _line("18.400", 120, 350, 210, 370),
        _line("Rp15.900", 120, 375, 235, 397),
    ]
    windows = make_price_windows(lines, 800, 800)
    target = next((w for w in windows if w["price_values"] == [18400, 15900]), None)
    assert target is not None
    assert target["local_product_name_suggestion"] == ""
    assert target["local_product_name_suggestion_status"].startswith("DIKARANTINA")
    assert target["local_product_name_suggestion_risk"] == "TINGGI"

    image = {
        "source_name": "Katalogpromosi.com",
        "campaign_title": "Promo Mingguan",
        "article_url": "https://example.test/a",
        "image_url": "https://example.test/p.jpg",
    }
    products = [{"product_text": "Nutribra Pro 900 g", "product_validation_status": "eligible"}]
    candidates = build_candidates(image, lines, 800, 800, products, products)
    target_candidate = next(c for c in candidates if c["normal_price_candidate"] == 18400 and c["promo_price_candidate"] == 15900)
    assert target_candidate["matched_product_name"] == ""
    assert target_candidate["catalog_sku_candidate_1"] == ""
    assert target_candidate["review_priority"] == "TINGGI"


def test_review_priority_is_triage_only_and_flags_ambiguous_discount_or_quantity():
    from promo_poster_ocr import _review_priority

    priority, reason = _review_priority({
        "local_product_name_suggestion_status": "DIKARANTINA — TIDAK ADA NAMA PRODUK AMAN",
        "local_product_name_suggestion_risk": "TINGGI",
        "discount_claim_source": "KLAIM_DISKON_AMBIGU",
        "purchase_condition": "AMBIGU: beberapa syarat jumlah terdeteksi (2, 5 pcs)",
        "ocr_confidence": 0.95,
        "normal_price_candidate": 18000,
        "promo_price_candidate": 15000,
    })
    assert priority == "TINGGI"
    assert "dikarantina" in reason.lower()
    assert "diskon ambigu" in reason.lower()
    assert "jumlah pembelian ambigu" in reason.lower()


def test_candidate_metrics_distinguish_raw_and_workbook_visible_name_suggestions():
    from promo_poster_ocr import _candidate_metrics

    candidates = [
        {"local_product_name_suggestion": "Sunlight Lime 910 ml", "matched_product_name": "Sunlight Lime 910 ml"},
        {"local_product_name_suggestion": "Nutribra Pro 900 g", "matched_product_name": ""},
    ]
    metrics = _candidate_metrics(candidates)
    assert metrics["candidate_local_name_suggestion_count"] == 2
    assert metrics["review_local_name_suggestion_count"] == 1


def test_image_summary_distinguishes_zero_candidates_from_preview_excluded_images():
    from promo_poster_ocr import _image_summary_rows, OCR_PIPELINE_VERSION

    url_in = "https://example.test/in.jpg"
    url_out = "https://example.test/out.jpg"
    url_stale = "https://example.test/stale.jpg"
    images = [
        {"image_url": url_in, "source_name": "Katalogpromosi.com", "campaign_title": "Promo"},
        {"image_url": url_out, "source_name": "Katalogpromosi.com", "campaign_title": "Promo"},
        {"image_url": url_stale, "source_name": "Katalogpromosi.com", "campaign_title": "Promo"},
    ]
    cache = {
        url_in: {"ocr_pipeline_version": OCR_PIPELINE_VERSION, "ocr_status": "ok", "lines_json": json.dumps([_line("no price", 1, 1, 20, 20)]), "ocr_backend": "rapidocr"},
        url_out: {"ocr_pipeline_version": OCR_PIPELINE_VERSION, "ocr_status": "ok", "lines_json": json.dumps([_line("Rp10.000", 1, 1, 20, 20)]), "ocr_backend": "rapidocr"},
        url_stale: {"ocr_pipeline_version": "old-version", "ocr_status": "ok", "lines_json": json.dumps([_line("Rp10.000", 1, 1, 20, 20)]), "ocr_backend": "rapidocr"},
    }
    rows = _image_summary_rows(images, cache, [], included_urls={url_in})
    by_url = {r["image_url"]: r for r in rows}
    assert by_url[url_in]["image_result_status"] == "OCR_ADA_TAPI_TIDAK_ADA_KANDIDAT_HARGA"
    assert by_url[url_out]["image_result_status"] == "TIDAK_TERMASUK_CAKUPAN_PREVIEW"
    assert by_url[url_stale]["image_result_status"] == "TIDAK_TERMASUK_CAKUPAN_PREVIEW"


def test_quarantined_identity_cannot_keep_prior_accepted_decision(tmp_path):
    from openpyxl import load_workbook
    from promo_poster_ocr import make_review_workbook

    output = tmp_path / "review.xlsx"
    candidate = {
        "candidate_id": "SAMEID123",
        "source_name": "Katalogpromosi.com",
        "campaign_title": "Promo HAP",
        "article_url": "https://example.test/a",
        "image_url": "https://example.test/p.jpg",
        "ocr_snippet": "Aplikasi gopay Bayar Pakai Pay | Rp18.400 | Rp15.900",
        "normal_price_candidate": 18400,
        "promo_price_candidate": 15900,
        "local_product_name_suggestion": "",
        "local_product_name_suggestion_raw": "Aplikasi gopay Bayar Pakai Pay",
        "local_product_name_suggestion_status": "DIKARANTINA — TIDAK ADA NAMA PRODUK AMAN",
        "local_product_name_suggestion_risk": "TINGGI",
        "local_product_name_quarantine_reason": "Saran nama didominasi teks pembayaran.",
        "payment_campaign_cue_count": 2,
        "review_priority": "TINGGI",
        "review_priority_reason": "Nama OCR dikarantina.",
    }
    prior = {
        "SAMEID123": {
            "ID Kandidat": "SAMEID123",
            "Keputusan": "TERIMA",
            "Nama Produk untuk Review": "Aplikasi gopay Bayar Pakai Pay",
            "Saran Nama Produk (OCR Lokal)": "Aplikasi gopay Bayar Pakai Pay",
            "Harga Normal Kandidat (Rp)": 18400,
            "Harga Promo Kandidat (Rp)": 15900,
            "Harga Normal untuk Review (Rp)": 18400,
            "Harga Promo untuk Review (Rp)": 15900,
        }
    }
    make_review_workbook([candidate], output, prior)
    wb = load_workbook(output, data_only=True)
    ws = wb["Kandidat OCR"]
    headers = [cell.value for cell in ws[1]]
    assert ws.cell(2, headers.index("Keputusan") + 1).value == "CEK"
    assert ws.cell(2, headers.index("Nama Produk untuk Review") + 1).value in (None, "")
    assert ws.cell(2, headers.index("Saran Nama Ditahan (OCR)") + 1).value == "Aplikasi gopay Bayar Pakai Pay"
    assert ws.cell(2, headers.index("Prioritas Review") + 1).value == "TINGGI"
    wb.close()


def test_regional_restriction_is_not_counted_as_payment_contamination():
    from promo_poster_ocr import make_price_windows, _contains_payment_campaign_cue

    assert not _contains_payment_campaign_cue("Tidak berlaku di Ambon")
    lines = [
        _line("Mie Seirocka Tori Cup 114g", 110, 210, 340, 238),
        _line("Tidak berlaku di Ambon", 110, 250, 350, 270),
        _line("Rp17.900", 120, 350, 220, 370),
        _line("Rp13.900", 120, 375, 230, 397),
    ]
    windows = make_price_windows(lines, 800, 800)
    target = next((w for w in windows if w["price_values"] == [17900, 13900]), None)
    assert target is not None
    assert target["local_product_name_suggestion"] == "Mie Seirocka Tori Cup 114g"
    assert target["payment_campaign_cue_count"] == 0
    assert target["local_product_name_suggestion_status"] == "DIPISAHKAN DARI TEKS PROMOSI/KETENTUAN"
    assert not target["local_product_name_suggestion_status"].startswith("DIKARANTINA")


def test_known_new_pack_callout_is_removed_without_dropping_product_name():
    from promo_poster_ocr import _clean_local_product_line

    cleaned = _clean_local_product_line("Baru!! Isi 5 Bungkus!! Mie Sedaap instant")
    assert "baru" not in cleaned.casefold()
    assert "bungkus" not in cleaned.casefold()
    assert "mie sedaap" in cleaned.casefold()


def test_fuzzy_payment_ocr_typo_is_detected_without_flagging_product_word_ovo_inside_novo():
    from promo_poster_ocr import _contains_payment_campaign_cue

    assert _contains_payment_campaign_cue("Bayear pdkai Aplikasi gopay")
    assert not _contains_payment_campaign_cue("Novo Shampoo 300ml")


def test_mixed_product_line_keeps_product_prefix_when_payment_suffix_is_present():
    from promo_poster_ocr import make_price_windows, _contains_payment_campaign_cue

    lines = [
        _line("AQUA Air Mnrl PET 600ml Aplikasi", 110, 210, 360, 238),
        _line("Bayar Pakai Aplikasi gopay", 110, 240, 360, 268),
        _line("11.900", 120, 350, 210, 370),
        _line("Rp10.900", 120, 375, 235, 397),
    ]
    windows = make_price_windows(lines, 800, 800)
    target = next((w for w in windows if w["price_values"] == [11900, 10900]), None)
    assert target is not None
    assert target["local_product_name_suggestion"] == "AQUA Air Mnrl PET 600ml"
    assert not _contains_payment_campaign_cue(target["local_product_name_suggestion"])
