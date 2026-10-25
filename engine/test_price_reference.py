from __future__ import annotations

import csv
import sys
from datetime import date
from pathlib import Path

ENGINE_DIR = Path(__file__).resolve().parent
ROOT_DIR = ENGINE_DIR.parent
sys.path.insert(0, str(ENGINE_DIR))

import enrich_existing_promo_catalog_multisource as enhancer
from promo_price_reference import REFERENCE_FIELDS, find_reference_price, load_reference_database, update_reference_database

TODAY = date(2026, 10, 24)


def _normal_observation(**overrides):
    row = {
        "merchant_name": "Indomaret",
        "product_text": "Susu UHT Full Cream 1 L",
        "normal_price": 20000,
        "promo_price": 15000,
        "price_basis": "offer_price",
        "normal_price_evidence_text": "Harga Normal Rp20.000",
        "promo_price_evidence_text": "Harga Promo Rp15.000",
        "product_validation_status": "eligible",
        "freshness_status": "fresh",
        "scraped_at": "2026-10-23T10:15:00+07:00",
        "source_name": "Hemat.id",
        "source_type": "aggregator",
        "source_url": "https://www.hemat.id/promo-susu-test/",
        "verification_status": "aggregator_only",
        "region_or_channel": "",
        "campaign_title": "Promo Indomaret 23-25 Oktober 2026",
        "retained_from_previous_run": False,
        "price_quality_flag": None,
    }
    row.update(overrides)
    return row


def _promo_without_normal(**overrides):
    row = {
        "product_text": "Susu UHT Full Cream 1 L",
        "original_product_text": "Susu UHT Full Cream 1 L",
        "merchant_name": "Indomaret",
        "campaign_title": "Promo Indomaret 23-25 Oktober 2026",
        "campaign_type": "Katalog Promo",
        "promo_status": "active",
        "period_certainty": "confirmed",
        "period_precision": "range",
        "valid_from": "2026-10-23",
        "valid_until": "2026-10-25",
        "promo_price": 15000,
        "normal_price": None,
        "normal_price_provenance": "not_available",
        "reference_normal_price": None,
        "reference_price_source_name": None,
        "reference_price_source_url": None,
        "reference_price_evidence_text": None,
        "reference_price_observed_at": None,
        "reference_price_age_days": None,
        "reference_price_match_status": "not_checked",
        "promo_price_evidence_text": "Harga Promo Rp15.000",
        "normal_price_evidence_text": None,
        "discount_evidence_text": None,
        "discount_percent": None,
        "discount_percent_basis": "not_available",
        "discount_percent_confidence": "unknown",
        "discount_percent_calculated": None,
        "discount_savings_rp": None,
        "estimated_bundle_discount_percent": None,
        "price_basis": "offer_price",
        "purchase_condition": None,
        "deal_type": None,
        "offer_text": "Harga Promo Rp15.000",
        "raw_line": "Susu UHT Full Cream 1 L: Harga Promo Rp15.000",
        "source_url": "https://katalogpromosi.com/promo-indomaret-test/",
        "source_name": "Katalogpromosi.com",
        "source_type": "aggregator",
        "verification_status": "aggregator_only",
        "source_verification_label": "Agregator — belum diverifikasi retailer",
        "region_or_channel": "",
        "scraped_at": "2026-10-24T10:00:00+07:00",
        "freshness_status": "fresh",
        "product_validation_status": "eligible",
        "product_review_reason": None,
        "product_rejection_reason": None,
        "price_quality_flag": None,
        "data_quality_flags": None,
        "source_discount_claim_percent": None,
    }
    row.update(overrides)
    return row


def _write_reference(path: Path, rows: list[dict]):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=REFERENCE_FIELDS)
        writer.writeheader()
        writer.writerows(rows)


def _reference_dict(**overrides):
    row = {
        "merchant_name": "Indomaret",
        "product_text": "Susu UHT Full Cream 1 L",
        "normal_price": "20000",
        "observed_at": "2026-10-23T10:15:00+07:00",
        "source_name": "Hemat.id",
        "source_url": "https://www.hemat.id/promo-susu-test/",
        "evidence_text": "Harga Normal Rp20.000",
        "record_origin": "observed_scrape",
        "verification_status": "aggregator_only",
        "region_or_channel": "",
    }
    row.update(overrides)
    return row


def test_explicit_fresh_normal_price_is_saved_to_reference_database(tmp_path):
    path = tmp_path / "data" / "product_price_reference.csv"
    summary = update_reference_database(path, [_normal_observation()], TODAY)
    records = load_reference_database(path)
    assert summary["observations_added"] == 1
    assert len(records) == 1
    assert records[0]["merchant_name"] == "Indomaret"
    assert records[0]["product_text"] == "Susu UHT Full Cream 1 L"
    assert records[0]["normal_price"] == "20000"
    assert records[0]["record_origin"] == "observed_scrape"


def test_database_update_is_idempotent_for_same_page_product_and_day(tmp_path):
    path = tmp_path / "data" / "product_price_reference.csv"
    update_reference_database(path, [_normal_observation()], TODAY)
    summary = update_reference_database(path, [_normal_observation()], TODAY)
    assert summary["observations_added"] == 0
    assert len(load_reference_database(path)) == 1


def test_database_does_not_learn_missing_stale_retained_or_unverified_prices(tmp_path):
    path = tmp_path / "data" / "product_price_reference.csv"
    rows = [
        _normal_observation(scraped_at="", freshness_status="historical_or_unverified"),
        _normal_observation(freshness_status="preserved_previous", retained_from_previous_run=True),
        _normal_observation(normal_price_evidence_text="Rp20.000"),
        _normal_observation(price_quality_flag="normal_price_ambiguous_unit"),
        _normal_observation(product_validation_status="review"),
    ]
    summary = update_reference_database(path, rows, TODAY)
    assert summary["observations_added"] == 0
    assert len(load_reference_database(path)) == 0


def test_price_reference_finds_only_exact_merchant_product_size_and_variant(tmp_path):
    path = tmp_path / "reference.csv"
    _write_reference(path, [_reference_dict()])
    refs = load_reference_database(path)
    match = find_reference_price(refs, "Susu UHT Full Cream 1 L", "Indomaret", TODAY)
    assert match["status"] == "matched_recent_exact"
    assert match["normal_price"] == 20000
    assert match["age_days"] == 1
    cross = find_reference_price(refs, "Susu UHT Full Cream 1 L", "Alfamart", TODAY)
    assert cross["status"] == "matched_recent_cross_merchant"
    assert cross["match_scope"] == "cross_merchant"
    assert cross["merchant_name"] == "Indomaret"
    assert find_reference_price(refs, "Susu UHT Full Cream 250 ml", "Indomaret", TODAY)["status"] == "no_exact_match"
    assert find_reference_price(refs, "Susu UHT Chocolate 1 L", "Indomaret", TODAY)["status"] == "no_exact_match"



def test_newer_normal_price_replaces_previous_price_for_same_product_merchant(tmp_path):
    path = tmp_path / "data" / "product_price_reference.csv"
    first = _normal_observation(normal_price=20000, normal_price_evidence_text="Harga Normal Rp20.000", scraped_at="2026-10-22T10:00:00+07:00")
    update_reference_database(path, [first], TODAY)
    newer = _normal_observation(normal_price=23000, normal_price_evidence_text="Harga Normal Rp23.000", scraped_at="2026-10-24T10:00:00+07:00", source_name="Scanharga.com", source_url="https://scanharga.com/new-price")
    summary = update_reference_database(path, [newer], TODAY)
    records = load_reference_database(path)
    assert len(records) == 1
    assert records[0]["normal_price"] == "23000"
    assert records[0]["source_name"] == "Scanharga.com"
    assert summary["observations_updated"] == 1



def test_same_merchant_reference_takes_priority_over_newer_other_merchant():
    refs = [
        _reference_dict(merchant_name="Indomaret", normal_price="20000", observed_at="2026-10-23T10:00:00+07:00"),
        _reference_dict(merchant_name="Alfamart", normal_price="25000", observed_at="2026-10-24T10:00:00+07:00", source_url="https://example.test/alfamart"),
    ]
    match = find_reference_price(refs, "Susu UHT Full Cream 1 L", "Indomaret", TODAY)
    assert match["status"] == "matched_recent_exact"
    assert match["match_scope"] == "same_merchant"
    assert match["merchant_name"] == "Indomaret"
    assert match["normal_price"] == 20000

def test_missing_same_merchant_price_falls_back_to_latest_other_merchant(tmp_path):
    path = tmp_path / "reference.csv"
    _write_reference(path, [_reference_dict(merchant_name="Indomaret", normal_price="20000")])
    row = _promo_without_normal(merchant_name="Alfamart")
    refs = load_reference_database(path)
    match = find_reference_price(refs, "Susu UHT Full Cream 1 L", "Alfamart", TODAY)
    assert match["status"] == "matched_recent_cross_merchant"
    summary = enhancer.apply_price_reference([row], refs, today=TODAY)
    assert row["reference_normal_price"] == 20000
    assert row["reference_price_merchant_name"] == "Indomaret"
    assert row["reference_price_match_scope"] == "cross_merchant"
    assert row["discount_percent"] == 25.0
    assert row["reference_price_match_status"] == "matched_cross_merchant_and_discount_calculated_from_reference"
    assert summary["cross_merchant_references_found"] == 1
    clean = enhancer.make_clean_row(row)
    assert clean["Merchant Acuan Referensi"] == "Indomaret"
    assert "merchant lain" in clean["Dasar Diskon"].casefold()

def test_missing_current_normal_price_calculates_discount_from_recent_reference(tmp_path):
    path = tmp_path / "reference.csv"
    _write_reference(path, [_reference_dict()])
    refs = load_reference_database(path)
    row = _promo_without_normal()
    summary = enhancer.apply_price_reference([row], refs, today=TODAY)
    assert row["normal_price"] is None  # never pretend this price appeared in the current promo text
    assert row["reference_normal_price"] == 20000
    assert row["reference_price_merchant_name"] == "Indomaret"
    assert row["reference_price_match_scope"] == "same_merchant"
    assert row["discount_percent"] == 25.0
    assert row["discount_percent_calculated"] == 25.0
    assert row["discount_savings_rp"] == 5000
    assert row["discount_percent_basis"] == "calculated_from_price_reference"
    assert row["discount_percent_confidence"] == "medium"
    assert row["reference_price_match_status"] == "matched_and_discount_calculated_from_reference"
    assert summary["reference_calculated_discounts"] == 1
    clean = enhancer.make_clean_row(row)
    assert clean["Harga Asli dari Sumber (Rp)"] is None
    assert clean["Harga Normal Referensi (Rp)"] == 20000
    assert clean["Diskon (%)"] == 25.0
    assert "referensi harga historis" in clean["Dasar Diskon"].casefold()


def test_price_reference_calculation_passes_gate_only_with_source_and_evidence(tmp_path):
    path = tmp_path / "reference.csv"
    _write_reference(path, [_reference_dict()])
    row = _promo_without_normal()
    enhancer.apply_price_reference([row], load_reference_database(path), today=TODAY)
    issues = enhancer.export_quality_gate_issues(row)
    assert issues == [], issues
    invalid = dict(row)
    invalid["reference_price_source_url"] = None
    issues = enhancer.export_quality_gate_issues(invalid)
    assert "reference_discount_without_valid_reference" in issues


def test_missing_normal_price_is_recovered_from_explicit_offer_text_and_discount_calculated():
    row = _promo_without_normal(
        offer_text="Harga Promo Rp15.000 dari harga normalnya Rp20.000",
        raw_line="Susu UHT: Harga Promo Rp15.000 dari harga normalnya Rp20.000",
    )
    summary = enhancer.apply_price_reference([row], [], today=TODAY)
    assert row["normal_price"] == 20000
    assert row["normal_price_provenance"] == "current_offer_evidence"
    assert row["normal_price_evidence_text"]
    assert row["discount_percent"] == 25.0
    assert row["discount_percent_basis"] == "calculated_from_unit_prices"
    assert row["discount_savings_rp"] == 5000
    assert summary["normal_prices_recovered_from_offer_text"] == 1
    assert summary["discounts_calculated_from_recovered_source_prices"] == 1


def test_explicit_normal_price_recovery_is_blocked_for_bundle_offer():
    row = _promo_without_normal(
        offer_text="Beli 3 lebih hemat Rp11.200 dari harga normal Rp12.600",
        raw_line="Beli 3 lebih hemat Rp11.200 dari harga normal Rp12.600",
        price_basis="bundle_offer_amount",
        purchase_condition="Beli 3",
    )
    enhancer.apply_price_reference([row], [], today=TODAY)
    assert row["normal_price"] is None
    assert row["discount_percent"] is None


def test_expired_reference_is_not_used_to_infer_discount(tmp_path):
    path = tmp_path / "reference.csv"
    _write_reference(path, [_reference_dict(observed_at="2026-08-01T10:00:00+07:00")])
    row = _promo_without_normal()
    summary = enhancer.apply_price_reference([row], load_reference_database(path), today=TODAY)
    assert row["reference_normal_price"] is None
    assert row["discount_percent"] is None
    assert row["reference_price_match_status"] == "reference_too_old"
    assert summary["stale_reference_not_used"] == 1


def test_conflicting_same_day_reference_prices_are_not_used(tmp_path):
    path = tmp_path / "reference.csv"
    _write_reference(path, [
        _reference_dict(normal_price="20000", source_url="https://example.test/a"),
        _reference_dict(normal_price="21000", evidence_text="Harga Normal Rp21.000", source_name="Katalogpromosi.com", source_url="https://example.test/b"),
    ])
    row = _promo_without_normal()
    summary = enhancer.apply_price_reference([row], load_reference_database(path), today=TODAY)
    assert row["reference_normal_price"] is None
    assert row["discount_percent"] is None
    assert row["reference_price_match_status"] == "conflicting_prices_same_day"
    assert summary["ambiguous_reference_not_used"] == 1


def test_bundles_and_conditional_offers_do_not_get_historical_sku_discount(tmp_path):
    path = tmp_path / "reference.csv"
    _write_reference(path, [_reference_dict()])
    refs = load_reference_database(path)
    for overrides, expected_status in [
        ({"offer_text": "Beli 1 Gratis 1, Harga Promo Rp15.000", "deal_type": "Beli Gratis"}, "bundle_not_comparable"),
        ({"offer_text": "Bayar pakai QRIS cashback Rp2.000, Harga Promo Rp15.000", "deal_type": "Cashback"}, "conditional_or_payment_offer_not_comparable"),
        ({"purchase_condition": "Khusus member", "offer_text": "Harga Promo Rp15.000"}, "purchase_condition_not_comparable"),
    ]:
        row = _promo_without_normal(**overrides)
        enhancer.apply_price_reference([row], refs, today=TODAY)
        assert row["reference_normal_price"] == 20000
        assert row["discount_percent"] is None
        assert row["discount_percent_basis"] == "not_available"
        assert row["reference_price_match_status"] == expected_status


def test_reference_db_persists_manual_verified_official_prices(tmp_path):
    path = tmp_path / "data" / "product_price_reference.csv"
    _write_reference(path, [_reference_dict(
        record_origin="manual_verified",
        verification_status="retailer_official",
        evidence_text="Harga Normal Rp20.000",
        source_name="Retailer resmi",
        source_url="https://retailer.example/product",
    )])
    summary = update_reference_database(path, [], TODAY)
    refs = load_reference_database(path)
    assert summary["records_after"] == 1
    match = find_reference_price(refs, "Susu UHT Full Cream 1 L", "Indomaret", TODAY)
    assert match["status"] == "matched_recent_exact"
    assert match["record_origin"] == "manual_verified"
    assert match["verification_status"] == "retailer_official"


def test_manual_reference_wins_over_aggregator_price_when_same_day(tmp_path):
    path = tmp_path / "reference.csv"
    _write_reference(path, [
        _reference_dict(normal_price="20000"),
        _reference_dict(normal_price="22000", evidence_text="Harga Normal Rp22.000", source_name="Retailer resmi", source_url="https://retailer.example/product", record_origin="manual_verified", verification_status="retailer_official"),
    ])
    match = find_reference_price(load_reference_database(path), "Susu UHT Full Cream 1 L", "Indomaret", TODAY)
    assert match["status"] == "matched_recent_exact"
    assert match["normal_price"] == 22000
    assert match["record_origin"] == "manual_verified"


def test_excel_exports_reference_database_and_clearly_labels_reference_discount(tmp_path, monkeypatch):
    from openpyxl import load_workbook

    db = tmp_path / "data" / "product_price_reference.csv"
    _write_reference(db, [_reference_dict()])
    monkeypatch.setattr(enhancer, "PRICE_REFERENCE_CSV", db)
    row = _promo_without_normal()
    enhancer.apply_price_reference([row], load_reference_database(db), today=TODAY)
    workbook_path = tmp_path / "catalog.xlsx"
    assert enhancer.write_excel_catalog(workbook_path, [row], [row], [], [], [], [], {"status": "passed"})
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        assert "Referensi Harga" in workbook.sheetnames
        active = workbook["Aktif"]
        headers = [cell.value for cell in next(active.iter_rows(min_row=1, max_row=1))]
        assert "Harga Normal Referensi (Rp)" in headers
        assert "Dasar Diskon" in headers
        assert "URL Referensi Harga" in headers
        assert headers[:7] == ["Nama Produk", "Kategori Produk", "Merchant", "Harga Asli (Rp)", "Harga Promo (Rp)", "Diskon (%)", "Basis Harga Asli"]
        assert headers[7:11] == ["Harga Asli dari Sumber (Rp)", "Harga Normal Referensi (Rp)", "Rentang Referensi Harga (Rp)", "Dasar Diskon"]
        assert headers.index("Merchant Acuan Referensi") < headers.index("Sumber Referensi Harga")
        values = list(active.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        row_values = dict(zip(headers, values))
        assert row_values["Harga Asli dari Sumber (Rp)"] is None
        assert row_values["Harga Normal Referensi (Rp)"] == 20000
        assert row_values["Diskon (%)"] == 25.0
        reference_sheet = workbook["Referensi Harga"]
        assert reference_sheet.max_row == 2
        assert reference_sheet.cell(2, 3).value == 20000
    finally:
        workbook.close()
