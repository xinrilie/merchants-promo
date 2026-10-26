"""Export product rows that still need a sourced manual market-price check.

This workflow does not generate or consume AI price estimates. It remains useful
for market comparisons after OCR has recovered the official price from a poster.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from promo_price_reference import identity_key
from promo_names import clean_product_name
from promo_poster_overrides import apply_reviewed_poster_prices, DEFAULT_OVERRIDES_CSV

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "_internal" / "promo_multisource_products_all.csv"
OUT = ROOT / "hasil" / "Daftar_Cari_Harga_Manual.xlsx"
CATALOG_CANDIDATES = (
    ROOT / "hasil" / "Katalog_Promo_Final_Ringkas.csv",
    ROOT / "hasil" / "Katalog_Promo_Final.csv",
)
GENERIC_NON_PRODUCT_NAME_RE = re.compile(
    r"^(?:produk(?:\s+(?:belum\s+teridentifikasi|pilihan|tertentu))?|item\s+pilihan|"
    r"potongan(?:\s+harga)?|harga\s+(?:promo|spesial|normal)|promo(?:si)?|diskon|"
    r"min(?:imum)?\.?\s*(?:belanja|pembelanjaan|transaksi)|syarat\s+belanja|"
    r"penawaran\s+(?:menarik|spesial))$", re.I,
)
HEADERS = [
    "ID Produk", "Nama Produk", "Merchant Promo", "Harga Promo (Rp)",
    "Kata Kunci Pencarian (copy)", "Harga Referensi (Rp) - isi manual",
    "Nama Sumber", "URL Halaman Harga", "Bukti Harga pada Halaman",
    "Tanggal Cek (YYYY-MM-DD)", "Jenis Referensi (pasar/normal)", "Status"
]


def _main_catalog_product_keys():
    for path in CATALOG_CANDIDATES:
        if not path.exists():
            continue
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.DictReader(handle)
                headers = set(reader.fieldnames or [])
                field = "Nama Produk" if "Nama Produk" in headers else ("product_text" if "product_text" in headers else None)
                if not field:
                    continue
                keys = {identity_key(row.get(field, "")) for row in reader if identity_key(row.get(field, ""))}
                if keys:
                    return keys, True
        except (OSError, csv.Error):
            continue
    return set(), False


def read_products():
    if not SOURCE.exists():
        raise SystemExit(f"Data produk belum ada: {SOURCE}\nJalankan scraping terlebih dahulu.")
    with SOURCE.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    poster_stats = apply_reviewed_poster_prices(rows, DEFAULT_OVERRIDES_CSV)
    out, seen_source_rows, seen_products = [], set(), set()
    catalog_keys, catalog_gate_enabled = _main_catalog_product_keys()
    stats = {
        "source_candidates": 0, "duplicate_products": 0, "excluded_not_in_catalog": 0,
        "excluded_not_eligible": 0, "excluded_generic": 0,
        "excluded_with_normal_price": 0, "catalog_gate_enabled": catalog_gate_enabled,
        "poster_ocr_rows_applied": poster_stats.get("rows_updated", 0),
    }
    for row in rows:
        name = str(row.get("product_text") or "").strip()
        promo = str(row.get("promo_price") or "").strip()
        merchant = str(row.get("merchant_name") or "").strip()
        if not name or not promo:
            continue
        key = (identity_key(name), merchant.casefold(), promo)
        if key in seen_source_rows:
            continue
        seen_source_rows.add(key)
        if str(row.get("normal_price") or "").strip():
            stats["excluded_with_normal_price"] += 1
            continue
        stats["source_candidates"] += 1
        status = str(row.get("product_validation_status") or "").strip().casefold()
        if status and status != "eligible":
            stats["excluded_not_eligible"] += 1
            continue
        clean_name = clean_product_name(name).strip().strip(" .:-–—")
        if not clean_name or GENERIC_NON_PRODUCT_NAME_RE.fullmatch(clean_name):
            stats["excluded_generic"] += 1
            continue
        if catalog_gate_enabled and identity_key(name) not in catalog_keys:
            stats["excluded_not_in_catalog"] += 1
            continue
        slot = (identity_key(name), merchant.casefold(), promo)
        if slot in seen_products:
            stats["duplicate_products"] += 1
            continue
        seen_products.add(slot)
        out.append({
            "ID Produk": f"{len(out)+1:05d}",
            "Nama Produk": name,
            "Merchant Promo": merchant,
            "Harga Promo (Rp)": promo,
            "Kata Kunci Pencarian (copy)": f'"{name}" harga',
            "Harga Referensi (Rp) - isi manual": "",
            "Nama Sumber": "",
            "URL Halaman Harga": "",
            "Bukti Harga pada Halaman": "",
            "Tanggal Cek (YYYY-MM-DD)": "",
            "Jenis Referensi (pasar/normal)": "pasar",
            "Status": "BELUM DICARI",
        })
    return out, stats


def main():
    products, stats = read_products()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    existing = {}
    if OUT.exists():
        try:
            old = load_workbook(OUT, read_only=True, data_only=True)
            if "Cari Harga" in old.sheetnames:
                ws_old = old["Cari Harga"]
                header = [cell.value for cell in ws_old[1]]
                for values in ws_old.iter_rows(min_row=2, values_only=True):
                    old_row = dict(zip(header, values))
                    key = (str(old_row.get("Nama Produk") or "").casefold(), str(old_row.get("Merchant Promo") or "").casefold(), str(old_row.get("Harga Promo (Rp)") or ""))
                    if key[0]:
                        existing[key] = old_row
            old.close()
        except Exception:
            pass
    merged = []
    seen = set()
    for row in products:
        key = (row["Nama Produk"].casefold(), row["Merchant Promo"].casefold(), str(row["Harga Promo (Rp)"]))
        if key in seen:
            continue
        seen.add(key)
        old_row = existing.get(key)
        if old_row and any(old_row.get(field) for field in (
            "Harga Referensi (Rp) - isi manual", "URL Halaman Harga", "Bukti Harga pada Halaman"
        )):
            row.update({header: old_row.get(header, row.get(header, "")) for header in HEADERS})
            complete = all(str(row.get(field) or "").strip() for field in (
                "Harga Referensi (Rp) - isi manual", "Nama Sumber", "URL Halaman Harga",
                "Bukti Harga pada Halaman", "Tanggal Cek (YYYY-MM-DD)"
            ))
            row["Status"] = "SIAP DIIMPOR" if complete else "PERIKSA ISIAN"
        merged.append(row)

    wb = Workbook()
    ws = wb.active
    ws.title = "Cari Harga"
    ws.append(HEADERS)
    for row in merged:
        ws.append([row.get(header, "") for header in HEADERS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    widths = [12, 38, 23, 18, 48, 24, 24, 48, 42, 20, 25, 18]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {2, 5, 7, 8, 9})

    helpws = wb.create_sheet("Petunjuk")
    help_text = [
        "CARA PAKAI DAFTAR CARI HARGA PASAR",
        "1. Daftar ini hanya untuk produk yang belum memiliki harga normal dari sumber promo/poster OCR yang sudah ditinjau.",
        "2. Salin Kata Kunci Pencarian ke mesin pencari, lalu cari SKU, varian, ukuran dan isi yang sama persis.",
        "3. Isi harga, nama toko, URL halaman produk langsung, bukti teks harga yang terlihat, dan tanggal cek.",
        "4. Gandakan baris bila ingin mengumpulkan harga dari beberapa toko/domain independen.",
        "5. Jenis normal hanya jika halaman secara eksplisit menyebut harga normal; jika tidak, pilih pasar.",
        "6. Simpan dan tutup Excel, kemudian jalankan: python engine\\import_manual_price_references.py hasil\\Daftar_Cari_Harga_Manual.xlsx",
        "7. Referensi pasar tidak mengubah harga normal resmi dan tidak dianggap terverifikasi hanya karena nama toko diisi.",
        "8. Estimasi AI tanpa sumber sudah dihapus. Harga poster OCR baru digunakan setelah keputusan TERIMA pada workbook Review_Harga_Poster_OCR.xlsx.",
    ]
    for line in help_text:
        helpws.append([line])
    helpws.column_dimensions["A"].width = 140
    for row in helpws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT)
    print(f"Daftar pencarian harga pasar dibuat: {OUT}")
    print(f"Produk tanpa harga normal untuk diteliti: {len(merged)}")
    print(f"Produk yang sudah punya harga normal disaring: {stats['excluded_with_normal_price']}")
    print(f"Produk di luar katalog utama dilewati: {stats['excluded_not_in_catalog']}")
    print(f"Produk non-eligible/review dilewati: {stats['excluded_not_eligible']}")
    print(f"Nama generik/non-SKU dilewati: {stats['excluded_generic']}")
    print(f"Baris poster OCR yang diterapkan untuk penyaringan: {stats['poster_ocr_rows_applied']}")
    print("Tidak ada antrean AI atau file prompt estimasi yang dibuat.")


if __name__ == "__main__":
    main()
