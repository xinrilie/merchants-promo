"""Regression checks for the local poster OCR v2 upgrade."""
from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw
from openpyxl import load_workbook

import promo_poster_ocr as ocr


def test_packaging_percentage_is_not_a_discount_claim():
    text = "ENERGEN COKELAT | 30% LEBIH BANYAK SEREAL | Rp28.900 | Rp20.900"
    result = ocr.analyze_discount_claim(text)
    assert result["claim"] is None
    assert result["source"] == "PERSENTASE_KEMASAN_DIKECUALIKAN"
    assert "30%" in result["ignored"]

    parsed = ocr.interpret_price_pair(
        "Harga normal Rp28.900 | Harga promo Rp20.900 | 30% LEBIH BANYAK SEREAL"
    )
    assert parsed["normal_price_candidate"] == 28900
    assert parsed["promo_price_candidate"] == 20900
    assert parsed["discount_claim_percent"] == ""
    assert parsed["discount_claim_source"] == "PERSENTASE_KEMASAN_DIKECUALIKAN"


def test_packaging_percentage_split_over_ocr_lines_is_excluded():
    result = ocr.analyze_discount_claim(
        "ENERGEN | 30% | LEBIH BANYAK SEREAL | Rp28.900 | Rp20.900"
    )
    assert result["claim"] is None
    assert "30%" in result["ignored"]


def test_explicit_discount_is_kept_and_conflicting_claims_are_ambiguous():
    result = ocr.analyze_discount_claim(
        "Diskon 30% | Harga normal Rp28.900 | Harga promo Rp20.900"
    )
    assert result["claim"] == 30.0
    assert result["source"] == "KATA_DISKON"

    result = ocr.analyze_discount_claim("Diskon 20% | Promo diskon 30% | Rp20.000")
    assert result["claim"] is None
    assert result["source"] == "KLAIM_DISKON_AMBIGU"


def test_context_does_not_mistake_packaging_extra_for_retailer_multibuy():
    context = ocr.classify_promo_context("KODOMO | BUY 1 EXTRA 1 | Rp12.900")
    assert not context["purchase_condition"]
    assert "TEKS TAMBAHAN/KLAIM KEMASAN TERDETEKSI" in context["promo_context_tags"]
    assert "PROMO MULTI-ITEM TERDETEKSI" not in context["promo_context_tags"]

    context = ocr.classify_promo_context("Beli 2 Gratis 1 | Member Rp15.000 | Harga spesial")
    assert context["purchase_condition"].casefold() == "beli 2 gratis 1"
    assert context["member_price_candidate"] == 15000
    assert "LABEL HARGA SPESIAL TERDETEKSI" in context["promo_context_tags"]


def test_stacked_price_pair_has_explicit_method_and_needs_review():
    lines = [
        {"text": "Rp28.900", "confidence": 0.9, "x0": 450, "x1": 560, "y0": 790, "y1": 820},
        {"text": "Rp20.900", "confidence": 0.9, "x0": 450, "x1": 560, "y0": 850, "y1": 880},
    ]
    parsed = ocr.interpret_price_pair("Rp28.900 | Rp20.900")
    pair = ocr._infer_stacked_price_pair(parsed, lines, 1200, 1500)
    assert pair["normal_price_candidate"] == 28900
    assert pair["promo_price_candidate"] == 20900
    assert pair["price_pair_method"] == "INFERENSI_GEOMETRI_SEJAJAR"


def test_color_price_tag_detector_finds_local_red_and_yellow_labels():
    image = Image.new("RGB", (1200, 1500), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((200, 510, 430, 555), fill=(255, 210, 0))  # yellow normal price
    draw.rectangle((200, 610, 455, 660), fill=(230, 0, 20))  # red promo price
    boxes = ocr._detect_color_price_tag_boxes(image)
    colors = {box[4] for box in boxes}
    assert "YELLOW" in colors
    assert "RED" in colors


def test_review_workbook_adds_audit_columns_without_changing_decision_gate(tmp_path: Path):
    candidate = {
        "candidate_id": "TEST-CANDIDATE", "source_name": "Poster lokal", "campaign_title": "Kampanye uji",
        "article_url": "https://example.test/article", "image_url": "https://example.test/poster.jpg",
        "ocr_snippet": "Rp28.900 | Rp20.900", "price_values_json": "[28900, 20900]",
        "normal_price_candidate": 28900, "promo_price_candidate": 20900,
        "discount_calculated_percent": 27.7, "discount_validation_status": "TIDAK ADA KLAIM DISKON",
        "discount_difference_pp": "", "price_pair_method": "INFERENSI_GEOMETRI_SEJAJAR",
        "promo_context_tags": "LABEL HARGA NORMAL/PROMO TERDETEKSI", "purchase_condition": "",
        "regional_scope": "", "member_price_candidate": "",
        "ignored_packaging_percentages": "30% — 30% LEBIH BANYAK SEREAL",
        "discount_claim_source": "PERSENTASE_KEMASAN_DIKECUALIKAN",
        "discount_claim_context": "30% | LEBIH BANYAK SEREAL",
    }
    out = tmp_path / "review.xlsx"
    ocr.make_review_workbook([candidate], out, {})
    wb = load_workbook(out, data_only=False)
    try:
        ws = wb["Kandidat OCR"]
        assert ws.max_row == 2
        assert ws.max_column == len(ocr.REVIEW_HEADERS) == 35
        assert ws.cell(2, 2).value == "CEK"
        assert ws.cell(2, 28).value == "INFERENSI_GEOMETRI_SEJAJAR"
        assert "30%" in ws.cell(2, 33).value
        assert ws.cell(2, 34).value == "PERSENTASE_KEMASAN_DIKECUALIKAN"
        assert "TERIMA" in " ".join(str(dv.formula1 or "") for dv in ws.data_validations.dataValidation)
    finally:
        wb.close()


def test_cache_and_line_schemas_version_the_pipeline_and_preserve_recognition_method():
    assert "ocr_pipeline_version" in ocr.CACHE_FIELDS
    assert ocr.OCR_PIPELINE_VERSION == "poster-ocr-v2.1-fast-crop"
    assert ocr.CANDIDATE_PIPELINE_VERSION == "poster-ocr-v2.2-spatial-context"
    assert "candidate_pipeline_version" in ocr.CANDIDATE_FIELDS
    assert "recognition_method" in ocr.LINE_FIELDS
    saved = {"lines_json": '[{"text":"20.900","recognition_method":"COLOR_PRICE_TAG_CROP_RED"}]'}
    lines = ocr._line_records(saved)
    assert lines[0]["recognition_method"] == "COLOR_PRICE_TAG_CROP_RED"


def test_color_crop_rescue_skips_tag_when_full_page_ocr_already_found_price(monkeypatch):
    image = Image.new("RGB", (1200, 1500), "white")
    boxes = [
        (200, 510, 430, 555, "YELLOW"),
        (600, 610, 830, 655, "RED"),
    ]
    monkeypatch.setattr(ocr, "_detect_color_price_tag_boxes", lambda _image: boxes)
    calls = []
    def fake_crop(_crop, _engine, _backend):
        calls.append(True)
        return "Rp20.900", 0.9
    monkeypatch.setattr(ocr, "_crop_ocr_price_text", fake_crop)
    lines = [{
        "text": "Rp28.900", "confidence": 0.9,
        "x0": 230, "y0": 520, "x1": 390, "y1": 545,
        "recognition_method": "FULL_PAGE",
    }]
    augmented = ocr._augment_with_color_price_tag_crops(image, lines, object(), "rapidocr")
    # The first tag is already covered by a full-page price. Only the missing
    # second tag receives a crop inference and is added as evidence.
    assert len(calls) == 1
    assert len(augmented) == 2
    assert any(line.get("recognition_method") == "COLOR_PRICE_TAG_CROP_RED" for line in augmented)


def test_color_crop_rescue_has_a_per_image_upper_bound(monkeypatch):
    image = Image.new("RGB", (1200, 1500), "white")
    boxes = [(100 + 100*i, 500 + 50*i, 180 + 100*i, 530 + 50*i, "RED") for i in range(12)]
    monkeypatch.setattr(ocr, "_detect_color_price_tag_boxes", lambda _image: boxes)
    calls = []
    def fake_crop(_crop, _engine, _backend):
        calls.append(True)
        return "not a price", 0.0
    monkeypatch.setattr(ocr, "_crop_ocr_price_text", fake_crop)
    ocr._augment_with_color_price_tag_crops(image, [], object(), "rapidocr")
    assert len(calls) == ocr.MAX_COLOR_PRICE_CROPS_PER_IMAGE == 6


def test_purchase_quantity_conditions_handle_ocr_suffixes_and_exclude_packaging_copy():
    two_piece = ocr._extract_purchase_condition("BELI 2 INDOFOOD | 16.000/2Pcs | R13.200 /2PCs")
    assert "Beli 2" in two_piece
    assert "2 pcs" in two_piece

    five_piece = ocr._extract_purchase_condition("BELI 5 INDOMIE | 16.500/5Pc | Rp15.300/5PcS")
    assert "Beli 5" in five_piece
    assert "5 pcs" in five_piece

    assert not ocr._extract_purchase_condition("KODOMO | BUY 1 EXTRA 1 | 30% LEBIH BANYAK SEREAL")


def test_spatial_price_clusters_do_not_merge_following_product_price():
    # Representative of the HAP poster: 27.500/17.900 are a stacked pair;
    # 33.900 and its Sakura/Kirei/Matsu title belong to the next product.
    lines = [
        {"text": "27.500", "confidence": .95, "x0": 76, "y0": 467, "x1": 195, "y1": 498},
        {"text": "17.900", "confidence": .92, "x0": 107, "y0": 496, "x1": 344, "y1": 575},
        {"text": "Var. Sakura, Kirei, Matsu", "confidence": .90, "x0": 79, "y0": 569, "x1": 326, "y1": 602},
        {"text": "725ml/725g Rp33.900", "confidence": .91, "x0": 79, "y0": 601, "x1": 326, "y1": 637},
    ]
    windows = ocr.make_price_windows(lines, 1273, 1800)
    value_sets = [row["price_values"] for row in windows]
    assert [27500, 17900] in value_sets
    assert [33900] in value_sets
    assert [27500, 17900, 33900] not in value_sets
    pair = next(row for row in windows if row["price_values"] == [27500, 17900])
    assert "33.900" not in pair["price_context"]
    assert "33.900" not in pair["product_context"]
    assert "33.900" not in pair["display_context"]


def test_spatial_name_context_does_not_attach_sambal_title_to_chil_kid_prices():
    lines = [
        {"text": "Var. Sambal Bangkok", "confidence": .9, "x0": 395, "y0": 963, "x1": 627, "y1": 1000},
        {"text": "Btl 135ml RP14.700/2Pcs", "confidence": .9, "x0": 375, "y0": 995, "x1": 627, "y1": 1031},
        {"text": "CHIL-KID 3DHA", "confidence": .95, "x0": 462, "y0": 1052, "x1": 731, "y1": 1099},
        {"text": "Madu, Vanila", "confidence": .94, "x0": 461, "y0": 1095, "x1": 647, "y1": 1135},
        {"text": "chilokid", "confidence": .88, "x0": 598, "y0": 1144, "x1": 755, "y1": 1188},
        {"text": "176.400", "confidence": .94, "x0": 460, "y0": 1245, "x1": 579, "y1": 1280},
        {"text": "168.500", "confidence": .93, "x0": 483, "y0": 1282, "x1": 707, "y1": 1367},
    ]
    windows = ocr.make_price_windows(lines, 1273, 1800)
    chilkid = next(row for row in windows if row["price_values"] == [176400, 168500])
    assert "CHIL" in chilkid["local_product_name_suggestion"].upper()
    assert "SAMBAL BANGKOK" not in chilkid["local_product_name_suggestion"].upper()
    sambal = next(row for row in windows if row["price_values"] == [14700])
    assert "SAMBAL BANGKOK" in sambal["local_product_name_suggestion"].upper()


def test_review_decision_migrates_across_context_changed_id_only_on_unique_pair(tmp_path: Path):
    old_row = {
        "ID Kandidat": "OLD-ID", "Keputusan": "TERIMA",
        "Nama Produk untuk Review": "Nama Produk Dikoreksi Manual",
        "Saran Nama Produk (OCR Lokal)": "Energen Cokelat 10 x 35g",
        "Harga Normal untuk Review (Rp)": 28900,
        "Harga Promo untuk Review (Rp)": 20900,
        "Harga Normal Kandidat (Rp)": 28900,
        "Harga Promo Kandidat (Rp)": 20900,
        "URL Gambar": "https://example.test/poster.jpg",
    }
    candidate = {
        "candidate_id": "NEW-ID", "source_name": "Test", "campaign_title": "Test",
        "article_url": "https://example.test/article", "image_url": "https://example.test/poster.jpg",
        "ocr_snippet": "different surrounding OCR context", "price_values_json": "[28900, 20900]",
        "normal_price_candidate": 28900, "promo_price_candidate": 20900,
        "local_product_name_suggestion": "Energen Cokelat 10 x 35g", "discount_calculated_percent": 27.7,
        "price_pair_method": "INFERENSI_GEOMETRI_SEJAJAR",
    }
    output = tmp_path / "migrated.xlsx"
    ocr.make_review_workbook([candidate], output, {"OLD-ID": old_row})
    wb = load_workbook(output, data_only=True)
    try:
        ws = wb["Kandidat OCR"]
        headers = [c.value for c in ws[1]]
        row = dict(zip(headers, next(ws.iter_rows(min_row=2, values_only=True))))
        assert row["ID Kandidat"] == "NEW-ID"
        assert row["Keputusan"] == "TERIMA"
        assert row["Nama Produk untuk Review"] == "Nama Produk Dikoreksi Manual"
        assert row["Harga Promo untuk Review (Rp)"] == 20900
    finally:
        wb.close()


def test_review_decision_does_not_migrate_when_fingerprint_is_ambiguous(tmp_path: Path):
    old_row = {
        "ID Kandidat": "OLD-ID", "Keputusan": "TERIMA",
        "Nama Produk untuk Review": "Manual", "Harga Normal Kandidat (Rp)": 28900,
        "Harga Promo Kandidat (Rp)": 20900, "URL Gambar": "https://example.test/poster.jpg",
    }
    candidates = []
    for new_id in ("NEW-A", "NEW-B"):
        candidates.append({
            "candidate_id": new_id, "source_name": "Test", "campaign_title": "Test",
            "article_url": "https://example.test/article", "image_url": "https://example.test/poster.jpg",
            "ocr_snippet": "same pair appears twice", "price_values_json": "[28900, 20900]",
            "normal_price_candidate": 28900, "promo_price_candidate": 20900,
            "discount_calculated_percent": 27.7,
        })
    output = tmp_path / "ambiguous.xlsx"
    ocr.make_review_workbook(candidates, output, {"OLD-ID": old_row})
    wb = load_workbook(output, data_only=True)
    try:
        ws = wb["Kandidat OCR"]
        assert [ws.cell(row, 2).value for row in (2, 3)] == ["CEK", "CEK"]
    finally:
        wb.close()


def test_review_approval_is_not_migrated_when_product_title_changes_across_spatial_fix(tmp_path: Path):
    old_row = {
        "ID Kandidat": "OLD-ID", "Keputusan": "TERIMA",
        "Nama Produk untuk Review": "Var. Sambal Bangkok",
        "Saran Nama Produk (OCR Lokal)": "Var. Sambal Bangkok",
        "Harga Normal Kandidat (Rp)": 176400, "Harga Promo Kandidat (Rp)": 168500,
        "URL Gambar": "https://example.test/hap.jpg",
    }
    candidate = {
        "candidate_id": "NEW-ID", "source_name": "Test", "campaign_title": "HAP Alfamidi",
        "article_url": "https://example.test/article", "image_url": "https://example.test/hap.jpg",
        "ocr_snippet": "CHIL-KID 3DHA | Rp176.400 | Rp168.500", "price_values_json": "[176400, 168500]",
        "normal_price_candidate": 176400, "promo_price_candidate": 168500,
        "local_product_name_suggestion": "CHIL-KID 3DHA Madu Vanila Box 780g",
        "discount_calculated_percent": 4.5,
    }
    output = tmp_path / "context_changed.xlsx"
    ocr.make_review_workbook([candidate], output, {"OLD-ID": old_row})
    wb = load_workbook(output, data_only=True)
    try:
        ws = wb["Kandidat OCR"]
        assert ws.cell(2, 2).value == "CEK"
        assert ws.cell(2, 3).value != "Var. Sambal Bangkok"
    finally:
        wb.close()
