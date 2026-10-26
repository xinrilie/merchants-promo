"""Import human-checked market price references from the manual Excel workbook.

AI price estimates are no longer accepted. This importer writes only references
with an explicit source URL and visible price evidence to the market database.
"""
from __future__ import annotations

import argparse
import csv
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook

from promo_market_prices import MARKET_REFERENCE_CSV, MARKET_REFERENCE_FIELDS, parse_idr_price, upsert_market_reference_database
from promo_names import product_name_size_signature, product_name_variant_signature
from promo_price_reference import identity_key

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "hasil" / "Daftar_Cari_Harga_Manual.xlsx"
DB = MARKET_REFERENCE_CSV
SEARCH_HOSTS = {
    "google.com", "www.google.com", "bing.com", "www.bing.com",
    "duckduckgo.com", "www.duckduckgo.com", "search.yahoo.com",
}


def price_int(value):
    n = parse_idr_price(value)
    return n if n is not None and 100 <= n <= 100_000_000 else None


def _normalize_input_row(row: dict, idx: int):
    name = str(row.get("Nama Produk") or "").strip()
    price = price_int(row.get("Harga Referensi (Rp) - isi manual"))
    source = str(row.get("Nama Sumber") or "").strip()
    url = str(row.get("URL Halaman Harga") or "").strip()
    evidence = str(row.get("Bukti Harga pada Halaman") or "").strip()
    observed = str(row.get("Tanggal Cek (YYYY-MM-DD)") or "").strip()[:10]
    kind = str(row.get("Jenis Referensi (pasar/normal)") or "pasar").strip().lower()
    title = str(row.get("Judul Halaman") or "").strip()

    parsed = urlparse(url)
    host = (parsed.hostname or "").lower().removeprefix("www.")
    issues = []
    if not name or not identity_key(name):
        issues.append("nama produk kosong/tidak valid")
    if price is None:
        issues.append("harga tidak valid (gunakan Rupiah, minimal Rp100)")
    if not source:
        issues.append("nama sumber kosong")
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        issues.append("URL harus berupa URL halaman harga langsung berawalan http(s)")
    elif (parsed.hostname.lower() in SEARCH_HOSTS or any(
        part in url.lower() for part in ("/search?", "/search/", "google.com/search", "bing.com/search")
    )):
        issues.append("URL hasil pencarian tidak boleh dipakai; gunakan halaman produk/harga langsung")
    if not evidence:
        issues.append("bukti harga pada halaman kosong")
    try:
        date.fromisoformat(observed)
    except ValueError:
        issues.append("tanggal harus YYYY-MM-DD")

    if issues:
        return None, f"Baris {idx}: " + "; ".join(issues)

    product_key = identity_key(name)
    normalized_kind = "normal" if kind in {"normal", "regular", "harga normal"} else "market"
    return ({
        "product_text": name,
        "product_key": product_key,
        "size_signature": "|".join(product_name_size_signature(name)),
        "variant_signature": "|".join(sorted(product_name_variant_signature(name))),
        "market_price": str(price),
        "observed_at": observed,
        "source_name": source,
        "source_domain": host,
        "source_url": url,
        "evidence_text": evidence,
        "page_title": title,
        "price_kind": normalized_kind,
        "source_kind": "manual_user_checked",
        "verification_method": "page_price_verified",
        "confidence": "manual_checked",
        "reference_key": "",
    }, None)


def _read_xlsx(path: Path):
    if not path.exists():
        raise SystemExit(f"File tidak ditemukan: {path}. Jalankan engine\\export_manual_price_search.py dahulu.")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Cari Harga" not in wb.sheetnames:
            raise SystemExit("Sheet 'Cari Harga' tidak ditemukan.")
        ws = wb["Cari Harga"]
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [str(value or "").strip() for value in header_row]
        entries, errors, total = [], [], 0
        required = {"Nama Produk", "Harga Referensi (Rp) - isi manual", "Nama Sumber", "URL Halaman Harga", "Bukti Harga pada Halaman", "Tanggal Cek (YYYY-MM-DD)"}
        missing = required - set(headers)
        if missing:
            raise SystemExit("Kolom Excel tidak lengkap: " + ", ".join(sorted(missing)))
        for idx, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            row = dict(zip(headers, values))
            if not any(str(row.get(k) or "").strip() for k in (
                "Harga Referensi (Rp) - isi manual", "Nama Sumber", "URL Halaman Harga", "Bukti Harga pada Halaman"
            )):
                continue
            total += 1
            normalized, error = _normalize_input_row(row, idx)
            if error:
                errors.append(error)
            else:
                entries.append(normalized)
        return entries, errors, total
    finally:
        wb.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description="Impor referensi harga pasar yang diperiksa manual dari Excel.")
    parser.add_argument("input", nargs="?", help="Workbook .xlsx. Jika kosong, gunakan hasil\\Daftar_Cari_Harga_Manual.xlsx.")
    args = parser.parse_args(argv)
    path = Path(args.input) if args.input else XLSX
    if not path.is_absolute():
        path = (Path.cwd() / path).resolve()
    if path.suffix.lower() == ".json":
        raise SystemExit("Impor JSON estimasi AI sudah dihapus. Gunakan poster OCR atau Excel berisi harga bersumber; JSON bukan lagi input yang diterima.")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise SystemExit("Format input belum didukung. Gunakan workbook .xlsx dari pencarian harga manual.")

    entries, errors, total = _read_xlsx(path)
    stats = upsert_market_reference_database(DB, entries)
    print(f"Input: {path}")
    print(f"Baris referensi dibaca: {total}")
    print(f"Referensi baru diimpor: {stats['records_added']}")
    print(f"Referensi diperbarui: {stats['records_updated']}")
    print(f"Baris ditolak karena tidak lengkap/tidak valid: {len(errors)}")
    for error in errors[:30]:
        print(" - " + error)
    if len(errors) > 30:
        print(f" - ... {len(errors) - 30} kesalahan lainnya")
    print(f"Database harga pasar: {DB}")
    print("Catatan: hanya referensi dengan URL, bukti harga, tanggal, dan sumber yang masuk; ini bukan otomatis harga normal resmi.")
    print("Langkah berikutnya: python engine\\enrich_existing_promo_catalog_multisource.py")
    return 0 if not errors else 2


if __name__ == "__main__":
    raise SystemExit(main())
