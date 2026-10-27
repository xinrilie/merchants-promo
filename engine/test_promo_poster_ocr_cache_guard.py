"""Regression tests for cache coverage and safe review-only rebuilds."""
from __future__ import annotations

import csv
import json

import pytest


def _write_csv(path, rows, fields):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _image(url):
    return {
        "image_url": url,
        "source_name": "Katalogpromosi.com",
        "campaign_title": "Campaign",
        "article_url": "https://example.test/campaign",
        "image_alt": "Poster",
        "period_status": "current",
    }


def _cache_row(ocr, url, *, status="ok", version=None, lines=None):
    if version is None:
        version = ocr.OCR_PIPELINE_VERSION
    if lines is None:
        lines = [{"text": "Produk X Rp10.000", "confidence": 0.95, "x0": 1, "y0": 1, "x1": 100, "y1": 20}]
    return {
        "image_url": url,
        "source_name": "Katalogpromosi.com",
        "campaign_title": "Campaign",
        "article_url": "https://example.test/campaign",
        "image_alt": "Poster",
        "period_status": "current",
        "scraped_at": "2026-10-26",
        "image_sha256": "testhash",
        "image_width": "800",
        "image_height": "1000",
        "ocr_status": status,
        "ocr_backend": "rapidocr",
        "ocr_pipeline_version": version,
        "ocr_error": "",
        "processed_at": "2026-10-26T18:00:00+07:00",
        "lines_json": json.dumps(lines),
    }


def test_cache_coverage_requires_current_valid_ocr_results(monkeypatch):
    import promo_poster_ocr as ocr

    images = [_image(f"https://example.test/{n}.jpg") for n in range(4)]
    cache = {
        images[0]["image_url"]: _cache_row(ocr, images[0]["image_url"]),
        images[1]["image_url"]: _cache_row(ocr, images[1]["image_url"], status="no_text", lines=[]),
        images[2]["image_url"]: _cache_row(ocr, images[2]["image_url"], version="poster-ocr-v1-old"),
        images[3]["image_url"]: _cache_row(ocr, images[3]["image_url"], status="failed", lines=[]),
    }

    usable, missing_or_stale = ocr._cache_coverage(images, cache)
    assert usable == 2
    assert missing_or_stale == 2


def test_full_review_rebuild_blocks_partial_cache_and_preserves_outputs(tmp_path, monkeypatch):
    import promo_poster_ocr as ocr

    image_index = tmp_path / "images.csv"
    products_csv = tmp_path / "products.csv"
    cache_csv = tmp_path / "cache.csv"
    lines_csv = tmp_path / "lines.csv"
    candidates_csv = tmp_path / "candidates.csv"
    review_xlsx = tmp_path / "review.xlsx"
    overrides_csv = tmp_path / "overrides.csv"

    images = [_image("https://example.test/cached.jpg"), _image("https://example.test/missing.jpg")]
    _write_csv(image_index, images, ["image_url", "source_name", "campaign_title", "article_url", "image_alt", "period_status"])
    _write_csv(products_csv, [], ["product_text", "source_url", "source_page_url", "product_validation_status", "merchant_name", "source_name"])
    _write_csv(cache_csv, [_cache_row(ocr, images[0]["image_url"])], ocr.CACHE_FIELDS)

    review_xlsx.write_bytes(b"existing main review workbook")
    lines_csv.write_text("existing lines", encoding="utf-8")
    candidates_csv.write_text("existing candidates", encoding="utf-8")
    overrides_csv.write_text("existing active overrides", encoding="utf-8")
    before = {
        review_xlsx: review_xlsx.read_bytes(),
        lines_csv: lines_csv.read_bytes(),
        candidates_csv: candidates_csv.read_bytes(),
        overrides_csv: overrides_csv.read_bytes(),
    }

    for name, value in {
        "IMAGE_INDEX": image_index,
        "PRODUCTS_CSV": products_csv,
        "CACHE_CSV": cache_csv,
        "LINES_CSV": lines_csv,
        "CANDIDATES_CSV": candidates_csv,
        "REVIEW_XLSX": review_xlsx,
        "OVERRIDES_CSV": overrides_csv,
    }.items():
        monkeypatch.setattr(ocr, name, value)
    monkeypatch.setattr(ocr, "import_reviewed_prices", lambda *a, **k: (_ for _ in ()).throw(AssertionError("must not import a partial review")))

    with pytest.raises(RuntimeError, match=r"1/2 gambar"):
        ocr.rebuild_review_from_cache()

    for path, content in before.items():
        assert path.read_bytes() == content


def test_limited_preview_uses_only_current_version_cache(tmp_path, monkeypatch):
    import promo_poster_ocr as ocr
    from openpyxl import Workbook, load_workbook

    image_index = tmp_path / "images.csv"
    products_csv = tmp_path / "products.csv"
    cache_csv = tmp_path / "cache.csv"
    lines_csv = tmp_path / "lines.csv"
    candidates_csv = tmp_path / "candidates.csv"
    review_xlsx = tmp_path / "review.xlsx"
    overrides_csv = tmp_path / "overrides.csv"

    current = _image("https://example.test/current.jpg")
    stale = _image("https://example.test/stale.jpg")
    _write_csv(image_index, [current, stale], ["image_url", "source_name", "campaign_title", "article_url", "image_alt", "period_status"])
    _write_csv(products_csv, [], ["product_text", "source_url", "source_page_url", "product_validation_status", "merchant_name", "source_name"])
    _write_csv(cache_csv, [
        _cache_row(ocr, current["image_url"]),
        _cache_row(ocr, stale["image_url"], version="poster-ocr-v1-old"),
    ], ocr.CACHE_FIELDS)

    for name, value in {
        "IMAGE_INDEX": image_index,
        "PRODUCTS_CSV": products_csv,
        "CACHE_CSV": cache_csv,
        "LINES_CSV": lines_csv,
        "CANDIDATES_CSV": candidates_csv,
        "REVIEW_XLSX": review_xlsx,
        "OVERRIDES_CSV": overrides_csv,
    }.items():
        monkeypatch.setattr(ocr, name, value)

    stats = ocr.rebuild_review_from_cache(limit=10)
    assert stats["images_indexed"] == 2
    assert stats["images_considered"] == 1
    assert stats["cache_usable_images"] == 1
    assert stats["cache_missing_or_stale_images"] == 1
    assert stats["cache_coverage_percent"] == 50.0
    assert stats["approved_override_count"] == "tidak disentuh (mode uji terpisah)"

    test_review = tmp_path / "Review_Harga_Poster_OCR_UjiTerbaru.xlsx"
    wb = load_workbook(test_review, data_only=True)
    ws = wb["Kandidat OCR"]
    assert ws.max_row >= 2
    for row in range(2, ws.max_row + 1):
        assert ws.cell(row, 17).value == current["image_url"]
    wb.close()


def test_incremental_ocr_with_partial_coverage_writes_preview_not_main_or_imports(tmp_path, monkeypatch):
    """An OCR batch must never replace the main workbook while cache coverage is partial."""
    import promo_poster_ocr as ocr

    image_index = tmp_path / "images.csv"
    products_csv = tmp_path / "products.csv"
    cache_csv = tmp_path / "cache.csv"
    lines_csv = tmp_path / "lines.csv"
    candidates_csv = tmp_path / "candidates.csv"
    review_xlsx = tmp_path / "Review_Harga_Poster_OCR.xlsx"
    overrides_csv = tmp_path / "overrides.csv"

    images = [_image(f"https://example.test/{n}.jpg") for n in range(3)]
    image_fields = ["image_url", "source_name", "campaign_title", "article_url", "image_alt", "period_status"]
    _write_csv(image_index, images, image_fields)
    _write_csv(products_csv, [], ["product_text", "source_url", "source_page_url", "product_validation_status", "merchant_name", "source_name"])
    _write_csv(cache_csv, [_cache_row(ocr, images[0]["image_url"])], ocr.CACHE_FIELDS)

    review_xlsx.write_bytes(b"preserve existing main workbook")
    overrides_csv.write_text("preserve existing active overrides", encoding="utf-8")
    original_review = review_xlsx.read_bytes()
    original_overrides = overrides_csv.read_bytes()
    written_review_paths = []

    for name, value in {
        "IMAGE_INDEX": image_index,
        "PRODUCTS_CSV": products_csv,
        "CACHE_CSV": cache_csv,
        "LINES_CSV": lines_csv,
        "CANDIDATES_CSV": candidates_csv,
        "REVIEW_XLSX": review_xlsx,
        "OVERRIDES_CSV": overrides_csv,
    }.items():
        monkeypatch.setattr(ocr, name, value)
    monkeypatch.setattr(ocr, "_load_image_index", lambda path=None: images)

    monkeypatch.setattr(ocr, "create_ocr_engine", lambda preference: ("rapidocr", object()))
    monkeypatch.setattr(ocr, "_download_image", lambda session, url: (b"fake-image", "testhash-new"))
    monkeypatch.setattr(ocr, "recognize_image", lambda body, engine, backend: ([
        {"text": "Produk Baru Rp10.000 Rp8.000", "confidence": 0.99, "x0": 10, "y0": 10, "x1": 200, "y1": 40}
    ], 800, 1000))
    monkeypatch.setattr(ocr, "_load_page_products", lambda path=products_csv: {})
    monkeypatch.setattr(ocr, "_load_catalog_products", lambda path=products_csv: [])
    monkeypatch.setattr(ocr, "_load_old_review", lambda path: {})
    monkeypatch.setattr(ocr.time, "sleep", lambda *args, **kwargs: None)

    def _fake_build_candidates(record, lines, width, height, product_rows, fallback_rows):
        candidate = {field: "" for field in ocr.CANDIDATE_FIELDS}
        candidate.update({
            "candidate_id": "cand-" + record["image_url"].rsplit("/", 1)[-1],
            "source_name": record["source_name"],
            "campaign_title": record["campaign_title"],
            "normal_price_candidate": 10000,
            "promo_price_candidate": 8000,
            "local_product_name_suggestion": "Produk Baru",
        })
        return [candidate]

    monkeypatch.setattr(ocr, "build_candidates", _fake_build_candidates)

    def _fake_make_review(candidates, path, prior=None):
        written_review_paths.append(path)
        path.write_bytes(b"preview workbook")

    monkeypatch.setattr(ocr, "make_review_workbook", _fake_make_review)
    monkeypatch.setattr(
        ocr,
        "import_reviewed_prices",
        lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("partial cache must not import review decisions")),
    )

    stats = ocr.run_ocr(limit=1, backend_preference="rapidocr")

    preview_path = tmp_path / "Review_Harga_Poster_OCR_UjiTerbaru.xlsx"
    assert stats["images_processed"] == 1
    assert stats["cache_usable_images"] == 2
    assert stats["cache_missing_or_stale_images"] == 1
    assert stats["cache_coverage_percent"] == pytest.approx(66.7, abs=0.1)
    assert stats["review_workbook"] == str(preview_path)
    assert stats["approved_override_count"] == "tidak disentuh (cakupan cache belum lengkap)"
    assert written_review_paths == [preview_path]
    assert preview_path.read_bytes() == b"preview workbook"
    assert review_xlsx.read_bytes() == original_review
    assert overrides_csv.read_bytes() == original_overrides
    assert stats["candidate_count"] == 2  # only the two current-version caches are included
