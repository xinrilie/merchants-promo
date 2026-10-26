from promo_poster_ocr import extract_prices, interpret_price_pair, make_price_windows, build_candidates


def test_reads_labelled_normal_promo_and_discount():
    result = interpret_price_pair("Harga Normal Rp31.500 | Harga Promo Rp17.500 | Diskon 44%")
    assert result["normal_price_candidate"] == 31500
    assert result["promo_price_candidate"] == 17500
    assert result["discount_claim_percent"] == 44.0
    assert result["discount_calculated_percent"] == 44.4


def test_does_not_pair_unlabelled_prices_even_when_spatial_mode_is_enabled():
    # Spatial adjacency alone is not enough to infer roles for the two numbers.
    result = interpret_price_pair("Rp100.000 | Rp89.000", allow_unlabelled_pair=True)
    assert result["normal_price_candidate"] == ""
    assert result["promo_price_candidate"] == ""
    assert "tanpa label" in result["candidate_note"] or "klaim diskon" in result["candidate_note"]


def test_unlabelled_pair_requires_a_matching_discount_claim():
    result = interpret_price_pair("Rp100.000 | Rp89.000 | Diskon 11%", allow_unlabelled_pair=True)
    assert result["normal_price_candidate"] == 100000
    assert result["promo_price_candidate"] == 89000
    assert result["discount_calculated_percent"] == 11.0


def test_one_labelled_price_does_not_guess_other_role_without_discount_claim():
    result = interpret_price_pair("Harga Normal Rp100.000 | Rp89.000")
    assert result["normal_price_candidate"] == 100000
    assert result["promo_price_candidate"] == ""


def test_does_not_pair_when_discount_math_disagrees():
    result = interpret_price_pair("Rp100.000 | Rp80.000 | Diskon 10%")
    assert result["normal_price_candidate"] == ""
    assert result["promo_price_candidate"] == ""
    assert "tidak cocok" in result["candidate_note"]


def test_extract_prices_skips_units_and_campaign_years():
    prices = extract_prices("Susu 1 L tahun 2026; Harga Promo Rp15.900")
    assert [row["price"] for row in prices] == [15900]


def test_candidate_preserves_source_poster_and_suggests_product_from_article():
    lines = [
        {"text": "Sunlight Lime 910 ml", "confidence": 0.95, "x0": 100, "y0": 120, "x1": 470, "y1": 165},
        {"text": "Harga Normal Rp20.000", "confidence": 0.95, "x0": 100, "y0": 180, "x1": 470, "y1": 220},
        {"text": "Harga Promo Rp16.900", "confidence": 0.95, "x0": 100, "y0": 230, "x1": 470, "y1": 270},
        {"text": "Diskon 15%", "confidence": 0.95, "x0": 100, "y0": 280, "x1": 470, "y1": 320},
    ]
    image = {"source_name": "Katalogpromosi.com", "campaign_title": "Promo Mingguan", "article_url": "https://example.test/article", "image_url": "https://example.test/poster.jpg"}
    products = [{"product_text": "Sunlight Lime 910 ml", "product_validation_status": "eligible"}]
    candidates = build_candidates(image, lines, 800, 1000, products)
    assert candidates
    assert all(row["article_url"] == image["article_url"] for row in candidates)
    assert all(row["image_url"] == image["image_url"] for row in candidates)
    assert any(row["normal_price_candidate"] == 20000 and row["promo_price_candidate"] == 16900 for row in candidates)


def test_article_url_match_is_normalized_for_www_and_trailing_slash(tmp_path):
    import csv
    from promo_poster_ocr import _load_page_products, _canonical_article_url, _safe_image_url

    path = tmp_path / "products.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["product_text", "source_url", "source_page_url", "product_validation_status"])
        writer.writeheader()
        writer.writerow({"product_text": "Sunlight Lime 910 ml", "source_url": "https://www.example.test/article/", "source_page_url": "", "product_validation_status": "eligible"})
    products = _load_page_products(path)
    assert products[_canonical_article_url("https://example.test/article")][0]["product_text"] == "Sunlight Lime 910 ml"
    assert not _safe_image_url("http://127.0.0.1/private.jpg")
    assert not _safe_image_url("http://192.168.1.2/private.jpg")
    assert _safe_image_url("https://images.example.test/poster.jpg")


def test_review_workbook_links_to_article_and_original_poster(tmp_path):
    from openpyxl import load_workbook
    from promo_poster_ocr import make_review_workbook

    output = tmp_path / "review.xlsx"
    candidate = {
        "candidate_id": "AB12CD34", "source_name": "Katalogpromosi.com",
        "campaign_title": "Promo Mingguan", "article_url": "https://example.test/article",
        "image_url": "https://example.test/poster.jpg", "ocr_snippet": "Harga Normal Rp20.000 Harga Promo Rp15.000",
        "matched_product_name": "Sunlight Lime 910 ml", "matched_product_score": 0.95,
        "match_note": "Ukuran cocok", "candidate_note": "Harga ditinjau manual",
        "normal_price_candidate": 20000, "promo_price_candidate": 15000,
        "discount_claim_percent": 25, "discount_calculated_percent": 25,
        "ocr_confidence": 0.94,
    }
    make_review_workbook([candidate], output)
    wb = load_workbook(output, read_only=False, data_only=False)
    ws = wb["Kandidat OCR"]
    assert ws["P2"].hyperlink.target == "https://example.test/article"
    assert ws["Q2"].hyperlink.target == "https://example.test/poster.jpg"
    wb.close()


def test_spatial_windows_keep_adjacent_product_tiles_separate():
    from promo_poster_ocr import make_price_windows

    # Two promo tiles share the same horizontal band but have independent SKU
    # names and price pairs. Their texts must never be combined into one pair.
    lines = [
        {"text": "Sunlight Lime 910 ml", "confidence": .97, "x0": 80, "y0": 100, "x1": 290, "y1": 140},
        {"text": "Harga Normal Rp20.000", "confidence": .96, "x0": 85, "y0": 160, "x1": 285, "y1": 195},
        {"text": "Harga Promo Rp16.900", "confidence": .96, "x0": 85, "y0": 205, "x1": 285, "y1": 240},
        {"text": "Selsun Shampoo Yellow 100 ml", "confidence": .97, "x0": 520, "y0": 100, "x1": 800, "y1": 140},
        {"text": "Harga Normal Rp48.500", "confidence": .96, "x0": 525, "y0": 160, "x1": 795, "y1": 195},
        {"text": "Harga Promo Rp40.900", "confidence": .96, "x0": 525, "y0": 205, "x1": 795, "y1": 240},
    ]
    windows = make_price_windows(lines, 1000, 1000)
    assert len(windows) == 2
    by_price = {tuple(row["price_values"]): row for row in windows}
    assert (20000, 16900) in by_price
    assert (48500, 40900) in by_price
    sunlight = by_price[(20000, 16900)]["context"]
    selsun = by_price[(48500, 40900)]["context"]
    assert "Sunlight" in sunlight and "Selsun" not in sunlight
    assert "Selsun" in selsun and "Sunlight" not in selsun



def test_catalog_match_is_left_blank_when_two_skus_score_ambiguous(monkeypatch):
    import promo_poster_ocr as ocr
    scores = {"Sunlight Lime 910 ml": (.82, "size okay"), "Sunlight Lemon 910 ml": (.79, "size okay")}
    monkeypatch.setattr(ocr, "_matching_score", lambda name, context: scores[name])
    name, score, note = ocr._suggest_product("Sunlight 910 ml", [
        {"product_text": "Sunlight Lime 910 ml", "product_validation_status": "eligible"},
        {"product_text": "Sunlight Lemon 910 ml", "product_validation_status": "eligible"},
    ])
    assert name == ""
    assert score == .82
    assert "ambigu" in note.lower() or "skor" in note.lower()


def test_catalog_match_deduplicates_duplicate_rows_for_same_sku(monkeypatch):
    import promo_poster_ocr as ocr
    monkeypatch.setattr(ocr, "_matching_score", lambda name, context: (.91, "size okay"))
    name, score, note = ocr._suggest_product("Sunlight Lime 910 ml", [
        {"product_text": "Sunlight Lime 910 ml", "product_validation_status": "eligible"},
        {"product_text": "Sunlight Lime 910 ml", "product_validation_status": "eligible"},
    ])
    assert name == "Sunlight Lime 910 ml"
    assert score == .91

def test_local_ocr_name_suggestion_is_available_when_catalog_match_is_missing():
    from promo_poster_ocr import make_price_windows, build_candidates

    lines = [
        {"text": "DUA BELIBIS Saus Tomat 535 ml", "confidence": .93, "x0": 110, "y0": 100, "x1": 360, "y1": 140},
        {"text": "Rp19.200", "confidence": .94, "x0": 120, "y0": 170, "x1": 235, "y1": 205},
        {"text": "Rp15.400", "confidence": .94, "x0": 120, "y0": 210, "x1": 235, "y1": 245},
    ]
    windows = make_price_windows(lines, 800, 800)
    assert windows and windows[0]["local_product_name_suggestion"]
    image = {"source_name": "Katalogpromosi.com", "campaign_title": "Promo Mingguan", "article_url": "https://example.test/a", "image_url": "https://example.test/p.jpg"}
    rows = build_candidates(image, lines, 800, 800, page_products=[])
    assert rows
    assert rows[0]["matched_product_name"] == ""
    assert rows[0]["local_product_name_suggestion"]
    assert "OCR lokal" in rows[0]["match_note"]



def test_local_ocr_suggestion_is_separate_from_catalog_matched_sku(tmp_path):
    from openpyxl import load_workbook
    from promo_poster_ocr import make_review_workbook

    output = tmp_path / "review.xlsx"
    candidate = {
        "candidate_id": "LOCAL123", "source_name": "Katalogpromosi.com",
        "campaign_title": "Promo", "article_url": "https://example.test/a",
        "image_url": "https://example.test/p.jpg", "ocr_snippet": "DUA BELIBIS Saus Tomat 535 ml Rp19.200 Rp15.400",
        "matched_product_name": "", "matched_product_score": .43,
        "local_product_name_suggestion": "DUA BELIBIS Saus Tomat 535 ml",
        "normal_price_candidate": "", "promo_price_candidate": "",
        "discount_claim_percent": "", "discount_calculated_percent": "",
        "ocr_confidence": .94, "match_note": "Tidak ada kecocokan katalog yang kuat",
        "candidate_note": "Dua harga tanpa label/klaim pendukung; tidak dipasangkan.",
    }
    make_review_workbook([candidate], output)
    wb = load_workbook(output, data_only=True)
    ws = wb["Kandidat OCR"]
    headers = [cell.value for cell in ws[1]]
    product_col = headers.index("Nama Produk untuk Review") + 1
    suggestion_col = headers.index("Saran Nama Produk (OCR Lokal)") + 1
    assert ws.cell(2, product_col).value in (None, "")
    assert ws.cell(2, suggestion_col).value == "DUA BELIBIS Saus Tomat 535 ml"
    wb.close()


def test_local_suggestion_with_two_distinct_sizes_is_not_combined_into_one_sku():
    from promo_poster_ocr import make_price_windows
    lines = [
        {"text": "Gandum Cokelat 140g Bourbon Choco 150g", "confidence": .92, "x0": 100, "y0": 100, "x1": 430, "y1": 140},
        {"text": "Rp20.000", "confidence": .95, "x0": 130, "y0": 170, "x1": 250, "y1": 200},
        {"text": "Rp18.000", "confidence": .95, "x0": 130, "y0": 205, "x1": 250, "y1": 235},
    ]
    windows = make_price_windows(lines, 800, 800)
    assert windows
    assert all("Gandum Cokelat 140g Bourbon Choco 150g" != row["local_product_name_suggestion"] for row in windows)


def test_rebuild_drops_old_unreviewed_guesses_that_were_auto_copied(tmp_path):
    from openpyxl import load_workbook
    from promo_poster_ocr import make_review_workbook

    output = tmp_path / "review.xlsx"
    candidate = {
        "candidate_id": "STALE123", "source_name": "Katalogpromosi.com",
        "campaign_title": "Promo", "article_url": "https://example.test/a",
        "image_url": "https://example.test/p.jpg", "ocr_snippet": "Rp42.100 | Rp37.900",
        "matched_product_name": "", "matched_product_score": 0.0,
        "local_product_name_suggestion": "Adult Diapers Pants L8",
        "normal_price_candidate": "", "promo_price_candidate": "",
        "discount_claim_percent": "", "discount_calculated_percent": "",
        "ocr_confidence": .92, "match_note": "SKU katalog tidak unik", "candidate_note": "Harga tanpa label; tidak dipasangkan.",
    }
    prior = {
        "STALE123": {
            "Keputusan": "CEK", "Nama Produk untuk Review": "Adult Diapers Pants L8",
            "Nama Produk Saran": "Adult Diapers Pants L8",
            "Harga Normal untuk Review (Rp)": 42100, "Harga Promo untuk Review (Rp)": 37900,
            "Harga Normal Kandidat (Rp)": 42100, "Harga Promo Kandidat (Rp)": 37900,
        }
    }
    make_review_workbook([candidate], output, prior)
    wb = load_workbook(output, data_only=True)
    ws = wb["Kandidat OCR"]
    assert ws["B2"].value == "CEK"
    assert ws["C2"].value in (None, "")
    assert ws["E2"].value in (None, "")
    assert ws["F2"].value in (None, "")
    assert ws["D2"].value == "Adult Diapers Pants L8"
    wb.close()

def test_review_workbook_labels_inputs_as_review_not_approval(tmp_path):
    from openpyxl import load_workbook
    from promo_poster_ocr import make_review_workbook

    output = tmp_path / "review.xlsx"
    candidate = {
        "candidate_id": "SAFE1234", "source_name": "Katalogpromosi.com",
        "campaign_title": "Promo", "article_url": "https://example.test/article",
        "image_url": "https://example.test/poster.jpg", "ocr_snippet": "Sunlight 910 ml Harga Normal Rp20.000 Harga Promo Rp16.900",
        "matched_product_name": "Sunlight Lime 910 ml", "matched_product_score": .94,
        "match_note": "SKU cocok", "candidate_note": "Tinjau poster",
        "normal_price_candidate": 20000, "promo_price_candidate": 16900,
        "discount_claim_percent": 15, "discount_calculated_percent": 15.5,
        "ocr_confidence": .96,
    }
    make_review_workbook([candidate], output)
    wb = load_workbook(output, data_only=True)
    ws = wb["Kandidat OCR"]
    headers = [cell.value for cell in ws[1]]
    assert ws["B2"].value == "CEK"
    assert "Harga Normal Disetujui (Rp)" not in headers
    assert "Harga Promo Disetujui (Rp)" not in headers
    assert ws["E2"].value == 20000
    assert ws["F2"].value == 16900
    assert "bukan persetujuan" in wb["Petunjuk"]["A3"].value.lower()
    wb.close()


def test_rebuild_review_from_cache_does_not_run_ocr_or_download(tmp_path, monkeypatch):
    import csv
    import json
    import promo_poster_ocr as ocr

    def write_csv(path, rows, fields):
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)

    image_url = "https://images.example.test/poster.jpg"
    article_url = "https://example.test/promo"
    image_index = tmp_path / "images.csv"
    products_csv = tmp_path / "products.csv"
    cache_csv = tmp_path / "cache.csv"
    lines_csv = tmp_path / "lines.csv"
    candidates_csv = tmp_path / "candidates.csv"
    review_xlsx = tmp_path / "review.xlsx"
    overrides_csv = tmp_path / "overrides.csv"

    write_csv(image_index, [{
        "image_url": image_url, "source_name": "Katalogpromosi.com", "campaign_title": "Promo Indomaret",
        "article_url": article_url, "image_alt": "Sunlight Lime promo", "period_status": "current",
    }], ["image_url", "source_name", "campaign_title", "article_url", "image_alt", "period_status"])
    write_csv(products_csv, [{
        "product_text": "Sunlight Lime 910 ml", "source_url": article_url, "source_page_url": article_url,
        "product_validation_status": "eligible", "merchant_name": "Indomaret", "source_name": "Katalogpromosi.com",
    }], ["product_text", "source_url", "source_page_url", "product_validation_status", "merchant_name", "source_name"])
    lines = [
        {"text": "Sunlight Lime 910 ml", "confidence": .97, "x0": 100, "y0": 100, "x1": 390, "y1": 140},
        {"text": "Harga Normal Rp20.000", "confidence": .97, "x0": 100, "y0": 160, "x1": 390, "y1": 200},
        {"text": "Harga Promo Rp16.900", "confidence": .97, "x0": 100, "y0": 205, "x1": 390, "y1": 245},
        {"text": "Diskon 15%", "confidence": .97, "x0": 100, "y0": 250, "x1": 390, "y1": 285},
    ]
    cache_row = {
        "image_url": image_url, "source_name": "Katalogpromosi.com", "campaign_title": "Promo Indomaret",
        "article_url": article_url, "image_alt": "Sunlight Lime promo", "period_status": "current",
        "scraped_at": "2026-10-26", "image_sha256": "fakehash", "image_width": "800", "image_height": "1000",
        "ocr_status": "ok", "ocr_backend": "rapidocr", "ocr_error": "", "processed_at": "2026-10-26T10:00:00+07:00",
        "lines_json": json.dumps(lines),
    }
    write_csv(cache_csv, [cache_row], ocr.CACHE_FIELDS)

    for name, value in {
        "IMAGE_INDEX": image_index, "PRODUCTS_CSV": products_csv, "CACHE_CSV": cache_csv,
        "LINES_CSV": lines_csv, "CANDIDATES_CSV": candidates_csv, "REVIEW_XLSX": review_xlsx,
        "OVERRIDES_CSV": overrides_csv,
    }.items():
        monkeypatch.setattr(ocr, name, value)

    def must_not_run(*args, **kwargs):
        raise AssertionError("OCR engine or downloader must not run in rebuild-review-only mode")

    monkeypatch.setattr(ocr, "create_ocr_engine", must_not_run)
    monkeypatch.setattr(ocr, "_download_image", must_not_run)
    stats = ocr.rebuild_review_from_cache()
    assert stats["images_processed"] == 0
    assert stats["ocr_backend"] == "cache-reuse"
    assert stats["candidate_count"] >= 1
    from openpyxl import load_workbook
    wb = load_workbook(review_xlsx, data_only=True)
    ws = wb["Kandidat OCR"]
    headers = [cell.value for cell in ws[1]]
    assert "Harga Normal untuk Review (Rp)" in headers
    assert "Harga Normal Disetujui (Rp)" not in headers
    assert ws["B2"].value == "CEK"
    wb.close()


def test_review_note_refreshes_legacy_auto_note_and_flags_discount_mismatch(tmp_path):
    from openpyxl import load_workbook
    from promo_poster_ocr import make_review_workbook

    output = tmp_path / "review.xlsx"
    candidate = {
        "candidate_id": "STALEDISK01", "source_name": "Katalogpromosi.com",
        "campaign_title": "Promo", "article_url": "https://example.test/a",
        "image_url": "https://example.test/p.jpg",
        "ocr_snippet": "Deluxe Kitchen Towel Harga Normal Rp47.900 Harga Promo Rp30.500 Diskon 34%",
        "matched_product_name": "Deluxe Kitchen Towel 3 x 80", "matched_product_score": .91,
        "match_note": "Kecocokan perlu konfirmasi", "candidate_note": "Pasangan harga lokal cocok dengan klaim diskon; tetap konfirmasi visual.",
        "normal_price_candidate": 47900, "promo_price_candidate": 30500,
        "discount_claim_percent": 34, "discount_calculated_percent": 36.3,
        "ocr_confidence": .97,
    }
    prior = {"STALEDISK01": {
        "Keputusan": "CEK",
        "Harga Normal Kandidat (Rp)": 47900,
        "Harga Promo Kandidat (Rp)": 30500,
        "Harga Normal untuk Review (Rp)": 47900,
        "Harga Promo untuk Review (Rp)": 30500,
        "Catatan Review": "Kandidat lokal: angka terbesar dianggap harga normal dan terkecil harga promo; WAJIB cocokkan poster sebelum TERIMA.",
    }}
    make_review_workbook([candidate], output, prior)
    wb = load_workbook(output, data_only=True)
    ws = wb["Kandidat OCR"]
    headers = [cell.value for cell in ws[1]]
    note = ws.cell(2, headers.index("Catatan Review") + 1).value
    status = ws.cell(2, headers.index("Validasi Klaim Diskon") + 1).value
    difference = ws.cell(2, headers.index("Selisih Diskon (poin persentase)") + 1).value
    assert "angka terbesar" not in note.lower()
    assert "REVIEW OCR AUTO:" in note
    assert "tidak cocok" in note.lower()
    assert status.startswith("TIDAK COCOK")
    assert difference == 2.3
    wb.close()


def test_review_workbook_adds_row_specific_sku_dropdown_choices(tmp_path):
    from openpyxl import load_workbook
    from promo_poster_ocr import make_review_workbook

    output = tmp_path / "review.xlsx"
    candidate = {
        "candidate_id": "SKU-DROPDOWN-1", "source_name": "Katalogpromosi.com",
        "campaign_title": "Promo", "article_url": "https://example.test/a",
        "image_url": "https://example.test/p.jpg", "ocr_snippet": "Sunlight Lime 910 ml Rp20.000 Rp16.900",
        "matched_product_name": "Sunlight Lime 910 ml", "matched_product_score": .91,
        "local_product_name_suggestion": "Sunlight Lime 910 ml", "match_note": "Ada kecocokan katalog",
        "candidate_note": "Periksa poster", "normal_price_candidate": "", "promo_price_candidate": "",
        "discount_claim_percent": 15, "discount_calculated_percent": "",
        "catalog_sku_candidate_1": "Sunlight Lime 910 ml", "catalog_sku_candidate_2": "Sunlight Lemon 910 ml",
        "catalog_sku_candidate_3": "Sunlight Extra Power 900 ml", "catalog_sku_score_1": .91,
        "catalog_sku_score_2": .78, "catalog_sku_score_3": .65,
    }
    make_review_workbook([candidate], output)
    wb = load_workbook(output, data_only=True)
    ws = wb["Kandidat OCR"]
    headers = [cell.value for cell in ws[1]]
    assert ws.cell(2, headers.index("Kandidat SKU 1") + 1).value == "Sunlight Lime 910 ml"
    assert ws.cell(2, headers.index("Kandidat SKU 2") + 1).value == "Sunlight Lemon 910 ml"
    validations = list(ws.data_validations.dataValidation)
    sku_validation = next(v for v in validations if "C2" in str(v.sqref))
    assert sku_validation.formula1 == "=$S2:$V2"
    status = ws.cell(2, headers.index("Validasi Klaim Diskon") + 1).value
    assert status == "BELUM BISA DIBANDINGKAN"
    wb.close()


def test_build_candidates_suggests_several_catalog_skus_without_auto_approving_them():
    from promo_poster_ocr import build_candidates

    lines = [
        {"text": "Sunlight Lime 910 ml", "confidence": .95, "x0": 100, "y0": 120, "x1": 470, "y1": 165},
        {"text": "Harga Normal Rp20.000", "confidence": .95, "x0": 100, "y0": 180, "x1": 470, "y1": 220},
        {"text": "Harga Promo Rp16.900", "confidence": .95, "x0": 100, "y0": 230, "x1": 470, "y1": 270},
        {"text": "Diskon 15%", "confidence": .95, "x0": 100, "y0": 280, "x1": 470, "y1": 320},
    ]
    image = {"source_name": "Katalogpromosi.com", "campaign_title": "Promo Mingguan", "article_url": "https://example.test/article", "image_url": "https://example.test/poster.jpg"}
    page_products = [{"product_text": "Sunlight Lime 910 ml", "product_validation_status": "eligible"}]
    fallback_products = [
        {"product_text": "Sunlight Lemon 910 ml", "product_validation_status": "eligible"},
        {"product_text": "Sunlight Extra Power 900 ml", "product_validation_status": "eligible"},
        {"product_text": "Sunlight Lime 910 ml", "product_validation_status": "eligible"},
    ]
    candidates = build_candidates(image, lines, 800, 1000, page_products, fallback_products)
    assert candidates
    assert any(row.get("catalog_sku_candidate_1") == "Sunlight Lime 910 ml" for row in candidates)
    row = candidates[0]
    assert "catalog_sku_score_1" in row
    assert row["matched_product_name"] == "Sunlight Lime 910 ml"
    # Choice fields are suggestions only; they don't alter review/approval state.
    assert "decision" not in row



def test_sku_match_rejects_size_only_and_generic_category_similarity():
    from promo_poster_ocr import _matching_score, _candidate_sku_choices

    cases = [
        ("Maestro Wijen Sangrai 100 g", "Kerupuk Tengiri 100g"),
        ("Maestro Wijen Sangrai 100 g", "Pasta Gigi MW 3 Whitening 100g"),
        ("Alfamart Facial Tissue Albi 50'S", "Fac Tissue Doraemon 70s"),
    ]
    for sku, context in cases:
        score, note = _matching_score(sku, context)
        assert score == 0.0, (sku, context, score, note)
        choices = _candidate_sku_choices(context, [{"product_text": sku, "product_validation_status": "eligible"}], [])
        assert all(name != sku for name, _ in choices), (sku, context, choices)


def test_sku_match_rejects_conflicting_flavour_or_package_size():
    from promo_poster_ocr import _matching_score

    score, note = _matching_score("Sunlight Lime 910 ml", "Sunlight Lemon 910 ml")
    assert score == 0.0
    assert "varian" in note.lower()

    score, note = _matching_score("Bebio Minyak Telon Plus 60 ml", "Bebio Minyak Telon Plus 100 ml")
    assert score == 0.0
    assert "ukuran" in note.lower()


def test_sku_match_keeps_exact_name_high_and_wrong_size_name_out():
    from promo_poster_ocr import _matching_score, _suggest_product

    score, note = _matching_score("Bebio Minyak Telon Plus 60 ml", "Bebio Minyak Telon Plus 60ml")
    assert score >= 0.95
    name, best_score, note = _suggest_product("Bebio Minyak Telon Plus 60ml", [
        {"product_text": "Bebio Minyak Telon Plus 100 ml", "product_validation_status": "eligible"},
        {"product_text": "Bebio Minyak Telon Plus 60 ml", "product_validation_status": "eligible"},
    ])
    assert name == "Bebio Minyak Telon Plus 60 ml"
    assert best_score >= 0.95


def test_no_match_option_is_in_sku_dropdown_and_importer_rejects_it(tmp_path):
    from openpyxl import load_workbook
    from promo_poster_ocr import make_review_workbook
    from promo_poster_overrides import NO_MATCH_SKU_SENTINEL, import_reviewed_prices

    review = tmp_path / "review.xlsx"
    overrides = tmp_path / "overrides.csv"
    candidate = {
        "candidate_id": "NOMATCH123456", "source_name": "Katalogpromosi.com",
        "campaign_title": "Promo", "article_url": "https://example.test/a",
        "image_url": "https://example.test/p.jpg", "ocr_snippet": "Kerupuk Tengiri Rp18.500 Rp13.900",
        "local_product_name_suggestion": "Kerupuk Tengiri 100g", "normal_price_candidate": 18500,
        "promo_price_candidate": 13900, "discount_claim_percent": 25,
        "discount_calculated_percent": 24.9, "ocr_confidence": .95,
        "catalog_sku_candidate_1": "", "catalog_sku_candidate_2": "", "catalog_sku_candidate_3": "",
    }
    make_review_workbook([candidate], review)
    wb = load_workbook(review)
    ws = wb["Kandidat OCR"]
    headers = [cell.value for cell in ws[1]]
    sku4_col = headers.index("Pilihan jika tidak ada SKU cocok") + 1
    assert ws.cell(2, sku4_col).value == NO_MATCH_SKU_SENTINEL
    validations = list(ws.data_validations.dataValidation)
    sku_validation = next(v for v in validations if "C2" in str(v.sqref))
    assert sku_validation.formula1 == "=$S2:$V2"
    # Selecting the sentinel and marking TERIMA must never import a fake product.
    ws.cell(2, headers.index("Keputusan") + 1, "TERIMA")
    ws.cell(2, headers.index("Nama Produk untuk Review") + 1, NO_MATCH_SKU_SENTINEL)
    ws.cell(2, headers.index("Harga Normal untuk Review (Rp)") + 1, 18500)
    ws.cell(2, headers.index("Harga Promo untuk Review (Rp)") + 1, 13900)
    wb.save(review)
    wb.close()

    stats = import_reviewed_prices(review, overrides)
    assert stats["accepted_or_updated"] == 0
    assert stats["rejected_invalid"] == 1
    assert any("bukan nama produk" in error for error in stats["errors"])


def test_sku_dropdown_does_not_recommend_conflicting_variant_even_with_same_brand_and_size():
    from promo_poster_ocr import _candidate_sku_choices, _suggest_product

    products = [
        {"product_text": "Sunlight Lime 910 ml", "product_validation_status": "eligible"},
        {"product_text": "Sunlight Lemon 910 ml", "product_validation_status": "eligible"},
    ]
    choices = _candidate_sku_choices("Sunlight Lime 910 ml", products, [], limit=3)
    assert choices
    assert choices[0][0] == "Sunlight Lime 910 ml"
    assert all(name != "Sunlight Lemon 910 ml" for name, _ in choices)
    matched, score, note = _suggest_product("Sunlight Lime 910 ml", products)
    assert matched == "Sunlight Lime 910 ml"
    assert score >= 0.95
