from datetime import date, datetime
from openpyxl import load_workbook

from promo_poster_ocr import make_review_workbook
from promo_poster_overrides import import_reviewed_prices, apply_reviewed_poster_prices
from promo_price_reference import load_reference_database, update_reference_database


def test_only_approved_poster_price_is_applied_to_matching_source_and_sku(tmp_path):
    review = tmp_path / "Review_Harga_Poster_OCR.xlsx"
    overrides = tmp_path / "overrides.csv"
    article = "https://example.test/promo-week"
    image = "https://images.example.test/poster.jpg"
    candidate = {
        "candidate_id": "ABCD1234EF5678",
        "source_name": "Katalogpromosi.com",
        "campaign_title": "Promo Mingguan",
        "article_url": article,
        "image_url": image,
        "image_alt": "Sunlight Lime 910 ml",
        "ocr_snippet": "Sunlight Lime 910 ml | Harga Normal Rp20.000 | Harga Promo Rp16.900 | Diskon 15%",
        "normal_price_candidate": 20000,
        "promo_price_candidate": 16900,
        "discount_claim_percent": 15,
        "discount_calculated_percent": 15.5,
        "ocr_confidence": 0.96,
        "matched_product_name": "Sunlight Lime 910 ml",
        "matched_product_score": 0.95,
        "match_note": "Ukuran cocok",
        "candidate_note": "Label harga cocok",
    }
    make_review_workbook([candidate], review)
    wb = load_workbook(review)
    ws = wb["Kandidat OCR"]
    ws["B2"] = "TERIMA"
    wb.save(review)
    wb.close()

    stats = import_reviewed_prices(review, overrides)
    assert stats["active_overrides"] == 1
    row = {
        "source_name": "Katalogpromosi.com",
        "source_type": "aggregator_product_listing",
        "source_url": article,
        "source_page_url": article,
        "product_text": "Sunlight Lime 910 ml",
        "merchant_name": "Indomaret",
        "promo_price": "16900",
        "normal_price": "",
        "offer_text": "Harga Promo Rp16.900",
        "price_basis": "offer_price",
        "scraped_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "freshness_status": "fresh",
        "product_validation_status": "eligible",
        "source_urls": article,
        "price_quality_flag": "normal_price_parse_failed_or_unverified",
    }
    applied = apply_reviewed_poster_prices([row], overrides)
    assert applied["rows_updated"] == 1
    assert row["normal_price"] == "20000"
    assert row["promo_price"] == "16900"
    assert row["normal_price_provenance"] == "poster_ocr_reviewed"
    assert row["price_parse_method"] == "poster_ocr_reviewed_manual"
    assert row["price_basis"] == "offer_price"
    assert "Harga Normal Rp20.000" in row["normal_price_evidence_text"]
    assert "normal_price_parse_failed_or_unverified" not in str(row.get("price_quality_flag") or "")

    # The approved OCR value becomes a source-backed normal-price reference and
    # uses the review timestamp, not the original scrape timestamp.
    db = tmp_path / "product_price_reference.csv"
    db_summary = update_reference_database(db, [row], today=date.today())
    stored = load_reference_database(db)
    assert db_summary["observations_added"] == 1
    assert len(stored) == 1
    assert stored[0]["record_origin"] == "reviewed_poster_ocr"
    assert stored[0]["observed_at"][:10] == str(row["poster_ocr_reviewed_at"])[:10]


def test_conflicting_existing_price_is_not_overwritten(tmp_path):
    review = tmp_path / "Review_Harga_Poster_OCR.xlsx"
    overrides = tmp_path / "overrides.csv"
    article = "https://example.test/promo-week"
    candidate = {
        "candidate_id": "CONFLICT1234567",
        "source_name": "Katalogpromosi.com", "campaign_title": "Promo", "article_url": article,
        "image_url": "https://images.example.test/poster.jpg", "image_alt": "Sunlight",
        "ocr_snippet": "Harga Normal Rp20.000 Harga Promo Rp16.900", "normal_price_candidate": 20000,
        "promo_price_candidate": 16900, "discount_claim_percent": 15, "discount_calculated_percent": 15.5,
        "ocr_confidence": 0.95, "matched_product_name": "Sunlight Lime 910 ml", "matched_product_score": 0.9,
        "match_note": "cocok", "candidate_note": "cek",
    }
    make_review_workbook([candidate], review)
    wb = load_workbook(review); wb["Kandidat OCR"]["B2"] = "TERIMA"; wb.save(review); wb.close()
    import_reviewed_prices(review, overrides)
    row = {"source_name": "Katalogpromosi.com", "source_url": article, "product_text": "Sunlight Lime 910 ml", "normal_price": "23000", "promo_price": "16900"}
    applied = apply_reviewed_poster_prices([row], overrides)
    assert row["normal_price"] == "23000"
    assert row.get("poster_ocr_override_conflict")
    assert applied["rows_conflicted"] == 1


def test_review_values_do_not_activate_until_decision_is_terima(tmp_path):
    review = tmp_path / "Review_Harga_Poster_OCR.xlsx"
    overrides = tmp_path / "overrides.csv"
    candidate = {
        "candidate_id": "CHECK123456789", "source_name": "Katalogpromosi.com",
        "campaign_title": "Promo", "article_url": "https://example.test/a",
        "image_url": "https://example.test/p.jpg", "ocr_snippet": "Produk Rp20.000 Rp16.900",
        "normal_price_candidate": 20000, "promo_price_candidate": 16900,
        "discount_claim_percent": "", "discount_calculated_percent": 15.5,
        "ocr_confidence": .95, "matched_product_name": "Sunlight Lime 910 ml",
        "matched_product_score": .9, "match_note": "Cek", "candidate_note": "Ambigu",
    }
    make_review_workbook([candidate], review)
    stats = import_reviewed_prices(review, overrides)
    assert stats["active_overrides"] == 0
    assert stats["accepted_or_updated"] == 0
    row = {
        "source_name": "Katalogpromosi.com", "source_url": "https://example.test/a",
        "product_text": "Sunlight Lime 910 ml", "normal_price": "", "promo_price": "",
    }
    apply_stats = apply_reviewed_poster_prices([row], overrides)
    assert apply_stats["rows_updated"] == 0
    assert row["normal_price"] == "" and row["promo_price"] == ""


def test_importer_still_accepts_legacy_approved_column_names(tmp_path):
    from openpyxl import Workbook

    review = tmp_path / "legacy-review.xlsx"
    overrides = tmp_path / "overrides.csv"
    wb = Workbook()
    ws = wb.active
    ws.title = "Kandidat OCR"
    headers = [
        "ID Kandidat", "Keputusan", "Nama Produk Disetujui", "Nama Produk Saran",
        "Harga Normal Disetujui (Rp)", "Harga Promo Disetujui (Rp)",
        "Harga Normal Kandidat (Rp)", "Harga Promo Kandidat (Rp)",
        "Klaim Diskon OCR (%)", "Diskon Hitung (%)", "Kecocokan Nama", "Keyakinan OCR",
        "Potongan Teks Poster OCR", "Sumber", "Kampanye", "URL Artikel", "URL Gambar", "Catatan Review",
    ]
    ws.append(headers)
    ws.append([
        "LEGACY12345678", "TERIMA", "Sunlight Lime 910 ml", "Sunlight", 20000, 16900,
        20000, 16900, 15, 15.5, "Cocok", .95, "Harga Normal Rp20.000 Harga Promo Rp16.900",
        "Katalogpromosi.com", "Promo", "https://example.test/a", "https://example.test/p.jpg", "OK",
    ])
    wb.save(review); wb.close()
    stats = import_reviewed_prices(review, overrides)
    assert stats["accepted_or_updated"] == 1
    assert stats["active_overrides"] == 1
